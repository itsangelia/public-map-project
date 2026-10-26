from pathlib import Path
from docx import Document
from docx.shared import Pt,Inches
import json,re,csv,zipfile
O=Path('/workspace/development-documentation-process');D=json.loads((O/'process-content.json').read_text())
def add(n,idx,en,zh):
 x=D[n-1];a=x['en'].split('\n\n');b=x['zh'].split('\n\n')
 if en in a[idx]:return
 a[idx]+=' '+en;b[idx]+=zh;x['en']='\n\n'.join(a);x['zh']='\n\n'.join(b)
x=D[4];x['en']=x['en'].replace('The phone panel retains a short session log, distinct from the temporary overhead bubble and from a permanent message archive.','The phone keeps up to 30 recent messages in the current page session; the database holds each connection’s latest speech state, not a permanent chat archive.');x['zh']=x['zh'].replace('手机面板保留简短的会话记录，与临时头顶气泡及永久消息档案有所区别。','手机保留当前页面会话中最多30条近期消息；数据库保留每条连接最近一次发言状态，而不是永久聊天档案。')
add(7,1,'Wishes remain available after their author leaves, separating contribution storage from temporary player presence.','作者离开后，愿望仍可保留，贡献内容的存储与临时玩家在场状态分开处理。')
add(9,1,'The counter supports one pair per room; a further visitor receives a busy message, and partner loss blocks further choices.','每个房间的柜台支持一组双人组合，额外访客会看到忙碌提示；搭档离开时，后续选择会被阻止。')
add(10,1,'Controls pause when partner readiness or fresh shared state is missing, rather than presenting a solo local round as continuing online play.','缺少搭档就绪或足够新的共享状态时，控制会暂停，而不会把单人的本地运行呈现为仍在继续的联机游戏。')
add(12,0,'Web Audio synthesises the tones rather than playing an uploaded piano recording.','声音由Web Audio合成，而不是播放上传的钢琴录音。')
D[13]['zh']=D[13]['zh'].replace('共同命名饮品与鱼','双方共同获得的有名字的饮品和鱼')
(O/'process-content.json').write_text(json.dumps(D,ensure_ascii=False,indent=2))
for kind,name in [('en','Development-Process-EN.docx'),('zh','Development-Process-ZH.docx'),('both','Development-Process-BILINGUAL.docx')]:
 doc=Document();doc.sections[0].top_margin=doc.sections[0].bottom_margin=Inches(.7);doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10.5)
 for i,x in enumerate(D):
  if i:doc.add_page_break()
  doc.add_heading(f"{x['page']:02d}  "+(x['title'] if kind=='en' else x['title_zh']),0)
  if kind=='both':doc.add_paragraph(x['title'],style='Subtitle')
  for en,zh in zip(x['en'].split('\n\n'),x['zh'].split('\n\n')):
   if kind in ['en','both']:doc.add_paragraph(en)
   if kind in ['zh','both']:doc.add_paragraph(zh)
  if i==13:doc.add_paragraph('Course sources: GRDE2053 Session 2, pp.6,8,10; Session 3, p.8; Session 5, slides6–11.',style='Caption')
 doc.save(O/name)
for both,name in [(False,'Development-Process-EN.md'),(True,'Development-Process-BILINGUAL.md')]:
 md='# Shared Campus Map — Development Documentation\n\n'
 for x in D:
  md+=f"## {x['page']:02d}. {x['title']}"+('｜'+x['title_zh'] if both else '')+'\n\n'
  for en,zh in zip(x['en'].split('\n\n'),x['zh'].split('\n\n')):md+=en+'\n\n'+(zh+'\n\n' if both else '')
 (O/name).write_text(md)
G=Document();G.add_heading('Development Documentation — 证据与排版安排',0)
G.add_paragraph('本版按开发阶段及实际迭代组织，不以A1对照作为主体。英文正文按15个页面主题安排，封面和参考文献另外排。每页需要真实图像/代码/记录支撑；中英对照用于核对，不是把两种语言全部塞进最终15页。')
for x in D:
 G.add_heading(f"{x['page']:02d} {x['title_zh']}",1);G.add_paragraph('证据安排：'+x['evidence']);G.add_paragraph(x['caption']);G.add_paragraph('依据：'+x['source'])
G.add_heading('证据使用边界',1)
for s in ['开发对话中的设备观察是实际过程证据，但不等于正式招募观众研究。没有记录的日期、人数、导师评价和复测结果不补写。','存储阻塞是有记录的复现问题，不是已确认的所有缺失玩家问题的唯一原因。','那次404没有明确根因诊断，不归因所有浏览器缓存；只写实际部署检查与后续版本措施。','祝福按匿名身份限制，不能写成永久真实个人只能一次。','控制输赢的是共享游戏状态，显示平滑不独立判定结果。','15页主题不是最终排版页数保证。配图后按教师正文最多15页检查，必要时压缩文字。','素材原作者、许可、课程规定引用格式、AI批准与声明，以及展场计时仍需按真实情况完成。']:G.add_paragraph('• '+s)
G.save(O/'Evidence-and-layout-guide-ZH.docx')
features=[('加入与共享地图','1–3','p5.js；Firebase匿名身份/RTDB/规则；HTTPS/QR；房间；昵称；三角色；独立连接；共享坐标与可见性；连接/人数；重连'),('移动与布局','4','手机跟随与边界；摇杆；地图铺满；35%放大/20%减速；昵称/气泡距离；活动打开停止移动与退出恢复'),('聊天','5','手机面板；文字；四快捷表情；最多30条当前会话消息；共享接收；5秒气泡；输入保护/发送反馈/冷却'),('互动入口','6','四入口素材；许愿树中央/英文标签/无手提示；饮品大小宽度位置；双鱼竿浅水；钢琴舞台/无木牌/+15%；靠近提示/点击/返回'),('许愿树','7–8','继承昵称；写/挂；共享存储与作者离开后保留；纸条；阅读昵称日期；计数/分页；预览/摆动；祝福/数量/身份限次；规则发布'),('饮品','9','双人方形匹配与姓名；四轮A/B/A/B；12原料选项/81配方；Your turn；已提交共享配方；视觉差异；成品放大/命名/返回；对齐修复；忙碌/搭档离开'),('钓鱼','10–11','双人匹配与中央倒计时；分轴控制；持续共享进度；显示平滑；20%初始/+10%/−4%；范围/3.5秒；0/100结果；涟漪/半浸浮标；四成功鱼/名称；失败Smug Finn；相同结果/重试/返回'),('钢琴','12','21键三排；单演奏者；WebAudio；开启/静音/距离音量；全地图听众；权限/iPhone恢复；地图内面板；退出间距；占用姓名通知/不显示听众键盘；图标动作音符；8秒实际音符续期；主动释放'),('整合发布','13','多人/消息/活动回归；规则与客户端匹配；Pages构建和来源；固定版本入口/匹配QR；预览到正式；保留备份；移除未使用路灯'),('评估与作者贡献','14–15','六原则/课堂理论；体验/氛围；短参与路径与展场检查；社会效果边界；AI用途/具体设计决定/引用')]
with (O/'Feature-and-process-coverage.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['功能与过程','正文页','覆盖细节']);w.writerows(features)
(O/'Sources-and-submission-notes.md').write_text('''# Sources and submission preparation\n\nThis is a15-topic development-process manuscript,not a complete image-laid-out assessment PDF. The supplied A2 brief requires visual process evidence and at most15 content pages,excluding cover/references.\n\n## Course and starting point\n- Zhao,A.(n.d.). Shared Campus Map:A Playful Shared Experience for Cross-Cultural Connection [Formal A1 exhibition concept;Social Interactive2.pdf],p.14 for the starting prototype.\n- GRDE2053 Social Interactive Experiences.(n.d.). A2_Brief,pp.2–5.\n- GRDE2053 Social Interactive Experiences.(n.d.). SIE–DCI A2:Exhibition Showcase marking rubric.\n- GRDE2053 Social Interactive Experiences.(n.d.). Session2,pp.6,8,10;Session3,p.8;Session5,slides6–11. Refer to course discussion rather than implying all primary theory publications have been read.\n\n## Implementation/evidence\n- Project source and development records:https://github.com/itsangelia/public-map-project . Relevant files and evidence locations are listed separately for each page in Evidence-and-layout-guide-ZH.docx.\n- p5.js documentation:https://p5js.org/\n- Firebase Realtime Database documentation:https://firebase.google.com/docs/database/\n- GitHub Pages documentation:https://docs.github.com/en/pages\n- MDN Web Audio API:https://developer.mozilla.org/en-US/docs/Web/API/Web_Audio_API\n\nUse the unit-outline citation style. Label and credit each figure;do not list only generic software websites as image sources. Original map,avatars,game message-board reference and supplied fish artwork require their own provenance. AI-generated/refined images require tool/use/adaptation records.\n\nStudent phone observations and development chat are not a recruited participant study. Do not invent dates or results for reconstructed evidence tables. The source records include stale early README limitations:current code and dated development entries take precedence when identifying completed features.\n\nTeacher PDFs and original A1 files are not re-published in this package. Tutor AI approval and media permissions are not claimed verified.\n''')
words=sum(len(re.findall(r"\b[\w’–-]+\b",x['en'])) for x in D)
(O/'READ-ME-ZH.md').write_text(f'''# Development Documentation — 完整开发过程稿\n\n本版替代此前过度围绕A1比较的写法，以你与助手的开发对话、源码和记录为依据，记录实际如何实现、遇到问题并反复修改。\n\n- Development-Process-BILINGUAL.docx：逐段完整中英文。\n- Development-Process-EN.docx / ZH.docx：单语言正文。\n- Evidence-and-layout-guide-ZH.docx：每页截图/代码/反馈与英文图注位置。\n- Feature-and-process-coverage.csv：功能与过程核对。\n\n15页主题，英文约{words}词；封面和参考文献另排，最终正文不能超过15页。它是正式开发文档文体，不是演讲稿。中英文对照用于阅读核对，最终排版不要默认双语全文仍能保持相同页数。\n\n正文依次讲基础/网络→扫码与玩家可见性→手机导航→聊天→入口美术→许愿贡献→发现与祝福→饮品→钓鱼同步难度→鱼素材与结果→钢琴→整合发布→六原则→评估与作者贡献。\n\n没有新做或编造现场测试；没有改变正式项目。待核实的研究/许可/批准事项放在来源和证据说明中。\n''')
# Whole-document checks.
assert len(D)==15
assert all(len(x['en'].split('\n\n'))==len(x['zh'].split('\n\n')) for x in D)
for kind,name in [('en','Development-Process-EN.docx'),('zh','Development-Process-ZH.docx'),('both','Development-Process-BILINGUAL.docx')]:
 paras=[x.text for x in Document(O/name).paragraphs]
 for x in D:
  if kind!='zh':
   for t in x['en'].split('\n\n'):assert t in paras
  if kind!='en':
   for t in x['zh'].split('\n\n'):assert t in paras
 assert zipfile.ZipFile(O/name).testzip() is None
body='\n'.join(x['en'] for x in D)
for s in ['in-memory authentication','acknowledged player write','30 recent messages','Anonymous','3.5','81','Lotus Fish','Purrmaid','eight-second','version','Sensory Feedback']:
 # Anonymous appears as lowercase in running prose.
 assert s.lower() in body.lower(),s
assert 'Token' not in body
verification={'topics':15,'english_words':words,'complete_bilingual_alignment':True,'source_implementation_reviewed':True,'new_real_device_tests_performed':False}
(O/'verification.json').write_text(json.dumps(verification,indent=2))
with zipfile.ZipFile(O/'Development-Process-package.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in O.iterdir():
  if f.suffix in ['.docx','.md','.csv','.json']:z.write(f,f.name)
print(verification)
