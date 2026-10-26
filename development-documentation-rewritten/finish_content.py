from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt
import json,zipfile,csv,re,shutil
O=Path('/workspace/development-documentation-rewritten');pages=json.loads((O/'page-content.json').read_text());pages[11]['title']='Reviewing the Finished Map'
(O/'page-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2))
checks=[('Observed problem','Revision / reported outcome'),('Only one player visible','Presence and joining reviewed;later device checks confirmed both players.'),('Small character and empty phone space','35% closer view and20% slower movement;layout approved during development.'),('Ingredient changes looked too similar','Clearer drink changes requested and reviewed within the selected style.'),('Fishing felt jerky and difficult','Movement refined;target tolerance,timing and progress loss adjusted separately.'),('Keys changed colour but iPhone was silent','Permissions/audio activation addressed;sound confirmed during device checks.')]
refs={0:'Initial concept: Social Interactive.pptx, slides 2, 5, 7–8. Confirm that this matches the submitted A1 version.',10:'Course sources: GRDE2053, Session 2, pp. 6, 8, 10; Session 3, p. 8; Session 5, slides 6–11. Theory is discussed here through the supplied course lectures.'}
for bilingual,name in [(False,'PDF-content-EN.docx'),(True,'PDF-content-EN-ZH.docx')]:
 d=Document();d.sections[0].top_margin=d.sections[0].bottom_margin=Inches(.7);d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10 if bilingual else 11)
 for i,p in enumerate(pages):
  if i:d.add_page_break()
  d.add_heading(f"{p['page']:02d}  {p['title']}",0)
  for para in p['en'].split('\n\n'):d.add_paragraph(para)
  if i==9:
   t=d.add_table(rows=1,cols=2);t.style='Light Shading Accent 1'
   for c,v in zip(t.rows[0].cells,checks[0]):c.text=v
   for r in checks[1:]:
    for c,v in zip(t.add_row().cells,r):c.text=v
  if i in refs:d.add_paragraph(refs[i],style='Caption')
  if bilingual:
   d.add_heading('中文说明与证据位置（排版时移出正文）',2);d.add_paragraph(p['zh']);d.add_paragraph('配图：'+p['figure']);d.add_paragraph('英文图注：'+p['caption'])
 d.save(O/name)
md='# Shared Campus — Development Documentation\n\n'
for p in pages:md+=f"## {p['page']:02d}. {p['title']}\n\n{p['en']}\n\n"
(O/'PDF-content-EN.md').write_text(md)
(O/'逐页中文与配图.md').write_text('# 中文理解与配图位置\n\n'+''.join(f"## 第{p['page']}页 {p['title']}\n\n{p['zh']}\n\n配图：{p['figure']}\n\n图注：{p['caption']}\n\n" for p in pages))
shutil.copy('/workspace/development-documentation/References-and-media-working-list.md',O/'References-and-media-working-list.md')
(O/'READ-ME-ZH.md').write_text('''# Shared Campus PDF逐页正文重写版\n\n13页主题，约2100英文词。封面和参考文献另外排，教师正文上限15页。\n\nPDF-content-EN.docx为正文，PDF-content-EN-ZH.docx含逐页中文解释、配图建议和图注。第10页已附开发检查小表，表中是开发记录与反馈，不是假设的用户研究。\n\n这套内容仍需你阅读确认，再用真实截图排成PDF。中文指导不要放进英文成品。没有制作新的测试照片或虚构测试日期。初始概念是否与正式A1一致、原始地图角色来源及许可、导师AI批准、unit outline引用格式仍需核实。第1页概念来源说明中的核对提示属于准备说明，确认后按实际情况移到备注或删除。第13页的待确认状态应在提交时更新为真实状态。\n\n配图要有编号、图注和来源；游戏参考图若使用，也要补原始出处。引用清单是工作表，不能当作全部媒体授权已核实的证明。正文没有声称完成正式观众研究，也没有证明改善跨文化归属感。\n\n这是排版文字，不是已经完成所有证据、引用和展览实测的最终提交PDF。正式地图没有变动。\n''')
# Content quality and packaging checks.
assert len(pages)==13
for p in pages:
 assert len(p['en'].split())>120
 assert 'User confirmations in the development conversation' not in p['en']
 assert 'The evidence should show' not in p['en']
principles=['Sensory Feedback','Embodied Interaction','Social Dynamics','Resilience & Motivation','Proactivity & Clarity','Ethics & Inclusion']
assert all(n in pages[10]['en'] for n in principles)
for name in ['PDF-content-EN.docx','PDF-content-EN-ZH.docx']:
 d=Document(O/name);texts=[p.text for p in d.paragraphs]
 for p in pages:
  for para in p['en'].split('\n\n'):assert para in texts
with zipfile.ZipFile(O/'PDF-page-writing-package.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in O.iterdir():
  if f.suffix in ['.docx','.md','.json']:z.write(f,f.name)
print('Checked13 page topics,both Word copies,and all six principle labels.')
