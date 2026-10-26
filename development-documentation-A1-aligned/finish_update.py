from pathlib import Path
from docx import Document
import json,csv,re,zipfile
O=Path('/workspace/development-documentation-A1-aligned');D=json.loads((O/'page-content.json').read_text())
notes=Document();notes.add_heading('依据正式A1更新：改动与证据',0)
notes.add_paragraph('核对文件：Social Interactive2.pdf，15页。封面为Shared Campus Map: A Playful Shared Experience for Cross-Cultural Connection，作者Angelia Zhao，A1 — Exhibition Concept。正文不再以先前Social Interactive.pptx概念草稿作为正式A1依据。')
notes.add_heading('必须更正的比较',1)
for text in ['正式A1第8页已经完成提示问答→个人手机→共享地图→合作活动的概念演变；不能把该转变归入A2。','正式A1没有列出代币系统。旧草稿中的代币想法不再作为正式A1到A2的删减。','正式A1案例为21 Swings与MMM Ball，不是21 Swings与Impulse。','A1第14页已实现单人p5.js昵称/三角色/确认/键盘移动，并列出五项未实现内容：QR连接、手机控制、实时联机、同时出现多个用户、共享活动。A2以实现这些内容为主要技术进展。','A1研究区分“已有朋友”与“跨越熟悉社交圈的机会不足”，因此不把国际学生描述为普遍没有朋友。','A1第15页已解释非人类角色的选择意义，以及Social Dynamics和Proactivity & Clarity的优先性。A2写保留与实现，不把它们说成新提出的原则。']:notes.add_paragraph(text)
notes.add_heading('A1到A2的对照',1)
rows=list(csv.reader((O/'A1-to-A2-comparison.csv').open(encoding='utf-8-sig')));t=notes.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,rows[0]):c.text=v
for r in rows[1:]:
 for c,v in zip(t.add_row().cells,r):c.text=v
notes.add_heading('逐页证据安排',1)
for p in D:
 notes.add_heading(f"{p['page']:02d} {p['title']}",2);notes.add_paragraph(p['zh']);notes.add_paragraph('配图：'+p['evidence']);notes.add_paragraph(p['caption']);notes.add_paragraph(p['source'])
notes.add_heading('仍需补齐，不能由A1文件自动确认',1)
notes.add_paragraph('实际开发截图、复测记录和展场计时；原始地图/角色/参考图来源与许可；unit outline引用格式；导师AI批准及正式声明。正式A1已确认，并不意味着其中引述的原始研究或图像授权也已独立核实。')
notes.save(O/'A1-alignment-and-evidence-guide-ZH.docx')
(O/'References-and-media-working-list.md').write_text('''# References and media credits — updated to the formal A1\n\nWorking entries. Format to the style required by the unit outline. No publication date is visible on the supplied A1 cover; n.d. is used rather than inventing one.\n\n- Zhao, A. (n.d.). Shared Campus Map: A Playful Shared Experience for Cross-Cultural Connection [A1 exhibition concept; supplied as Social Interactive2.pdf]. Relevant pages:2–4,5–6,8–15.\n- GRDE2053 Social Interactive Experiences. (n.d.). A2_Brief [Supplied assessment brief], pp.2–5.\n- GRDE2053 Social Interactive Experiences. (n.d.). SIE–DCI A2: Exhibition Showcase [Supplied marking rubric].\n- GRDE2053 Social Interactive Experiences. (n.d.). Session 2: Designing Participation [Lecture slides], pp.6,8,10.\n- GRDE2053 Social Interactive Experiences. (n.d.). Session 3: Designing with Care [Lecture slides], p.8 and Ethics & Inclusion discussion.\n- GRDE2053 Social Interactive Experiences. (n.d.). Session 5: Designing Interactions [Lecture slides], slides6–11.\n- p5.js. Documentation. https://p5js.org/\n- Firebase. Realtime Database documentation. https://firebase.google.com/docs/database/\n- GitHub. GitHub Pages documentation. https://docs.github.com/en/pages\n- MDN Web Docs. Web Audio API. https://developer.mozilla.org/en-US/docs/Web/API/Web_Audio_API\n- Project source and controlled interface captures: https://github.com/itsangelia/public-map-project . Record exact file/commit and capture date where used.\n\n## Research and precedent evidence\n\nThe formal A1 discusses research by Crawford, Sinanan and Gomes (2020), and McKenzie and Baldassar (2017). A2 currently refers to the A1 research synthesis; it does not pretend these primary publications have been independently read during this update. Consult original records before adding direct quotations or full primary-source claims.\n\nThe submitted A1 case studies are21 Swings /21 Balançoires (Daily tous les jours),p.5,and MMM Ball,p.6. Impulse appeared in an earlier draft and is not treated as a submitted A1 case study. Complete original project/image credits if reproducing their images. MMM Ball's reported participant count is not a test result from Shared Campus Map.\n\n## Media and AI items still requiring confirmation\n\nOriginal map/avatar/other supplied image creators and licences; source of the game message-board reference; individual generated/refined activity images; exact tool/asset records; tutor AI approval and formal declaration. The A1 author name does not establish original authorship of every reproduced image.\n\nFigure numbers,captions and source references are needed on the content pages as well as in the end references. Controlled screenshots are not participant-study photographs.\n''')
(O/'Presentation-A1-correction.md').write_text('''# Correction for the earlier presentation's concept-comparison paragraph\n\nThis update focuses on the13-page development-document text. Earlier presentation files retain their historical versions; if using their initial-concept comparison, replace references to token removal and the prompt-to-map transition with the paragraph below. Those statements were based on an earlier draft,not the submitted A1.\n\nEnglish replacement:\n\nMy submitted A1 had already developed the idea into an explorable map with avatars and shared play. Its p5.js prototype supported nickname entry,three avatar choices and keyboard movement,but not phone control or real-time multiplayer. A2 focused on implementing that shared experience through QR entry,independent players and specific activities,then refining it through device feedback.\n\n中文理解：正式A1已经确定共享地图和合作玩法方向，也有单人原型。A2是把手机联机和具体互动真正实现，再依据设备反馈完善；不能把概念稿较早阶段的转向当作A2新变化。\n''')
(O/'READ-ME-ZH.md').write_text('''# 依据正式A1更新的开发PDF内容\n\n已读Social Interactive2.pdf全15页，并检查封面及关键视觉页面。前一版依据早期PPT作出的阶段判断已更正。\n\n- Development-content-EN.docx：13页主题英文正文，正式文体，约2180词。\n- Development-content-EN-ZH.docx：英文＋中文说明＋配图/图注/来源。\n- A1-alignment-and-evidence-guide-ZH.docx：关键更正、A1/A2对照和逐页证据。\n- A1-to-A2-comparison.csv：可编辑对照表。\n- References-and-media-working-list.md：正式A1及相关资料工作引用。\n- Presentation-A1-correction.md：此前演讲稿“概念变化”段落的替换文案；本次未重排旧PPT。\n\n本版主线：从A1单人p5.js原型到A2多设备共享体验。自由探索、QR/手机的构想、合作活动方向和粉彩游戏风均在A1已有。A2写实现与测试，不重复提案。\n\n封面和参考文献另排；正文上限15页，本版13主题。需要加入真实过程截图、图注/引用并检查实际排版页数。A1已确认为正式版本，相关未知标记已移除；素材许可、导师AI批准、正式观众研究和展场计时仍不能假设已完成。\n\n本次没有修改已发布地图，也未将用户上传的原始A1/教师PDF公开转载。\n''')
# Verify exact delivery text and avoid historical A1 assumptions in the current body.
assert len(D)==13
body='\n'.join(x['en'] for x in D)
for false in ['token','cultural prompts','Impulse','relationship to the formal A1 submission requires confirmation']:assert false.lower() not in body.lower(),false
assert 'single-user' in D[0]['en'] and 'MMM Ball' in D[1]['en']
assert 'non-human avatar' in D[10]['en']
assert 'familiar social circles' in D[0]['en']
for name in ['Development-content-EN.docx','Development-content-EN-ZH.docx']:
 doc=Document(O/name);paras=[p.text for p in doc.paragraphs]
 for p in D:
  for s in p['en'].split('\n\n'):assert s in paras
 assert zipfile.ZipFile(O/name).testzip() is None
with zipfile.ZipFile(O/'Development-A1-aligned-package.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in O.iterdir():
  if f.suffix in ['.docx','.md','.csv','.json']:z.write(f,f.name)
print('Formal A1 baseline and13-page text checked;historical assumptions removed.')
