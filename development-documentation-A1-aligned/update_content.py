from pathlib import Path
import json,re,zipfile,csv
from docx import Document
from docx.shared import Inches,Pt
O=Path('/workspace/development-documentation-A1-aligned')
D=json.loads(Path('/workspace/development-documentation-academic/page-content.json').read_text())
def update(n,**fields):D[n-1].update(fields)
update(1,title='From the A1 Prototype to a Shared Experience',en='''The final A1 proposal established Shared Campus Map as a low-pressure opportunity for interaction among international and culturally diverse university students. Its research synthesis identified limited opportunities to engage beyond familiar social circles, rather than assuming that international students lacked friendships. Shared play was proposed as a reason for unfamiliar participants to notice and respond to one another.

A1 had already developed the concept from prompt-based participation into an explorable map with avatars and shared activities. It also included a single-user p5.js prototype testing nickname entry, three avatar choices, confirmation and keyboard movement. However, QR connection, phone control, real-time multiplayer and shared activities remained unimplemented.

A2 therefore concentrated on realising this proposed experience across separate devices. The existing prototype developed into a connected map with chat, shared wishes, collaborative drink-making, cooperative fishing and a shared piano. The principal progression was from a locally playable prototype to a shared system in which participants could see, respond to and influence one another.''',
zh='以正式A1为基准：社会问题是跨越熟悉圈子的互动机会有限，不是“国际学生没有朋友”。提示到地图的概念转变已在A1完成；A2的核心进展是单人原型到真实多人共享体验。删除代币与A2改变文化提示方向的旧表述。',
evidence='A1第14页单人原型与最终多人地图并排；小表列A1已实现和A2实现。A1第8页可作概念演变的小图，不标成A2阶段。',
caption='Figure 1. Progression from the single-user A1 prototype to a multi-device experience with shared player presence and activities.',
source='Zhao, A. (n.d.). Shared Campus Map: A Playful Shared Experience for Cross-Cultural Connection [A1 exhibition concept], pp. 2–4, 8, 10, 14. Supplied as Social Interactive2.pdf.')
update(2,title='Extending the A1 Technical Prototype',en='''A1’s technical exploration tested a single-user interaction flow: entering a nickname, selecting one of three avatars, confirming entry and moving through the map with arrow keys. This established the local interaction before the multiplayer system was developed. A2 retained the p5.js foundation while extending input and shared data across devices.

The phone now follows its own character, while the public display provides an overview of the campus. Firebase Realtime Database synchronises player positions and activity data, and GitHub Pages provides the web entry point used by the QR code. This implements the public-screen and personal-phone arrangement proposed in A1, informed by its analysis of MMM Ball’s browser-based participation.

Activities introduced further synchronisation requirements. Drink-making needs a common turn sequence and recipe; fishing combines separate directional inputs; the piano requires a single performer state. Messages and wishes must also reach other participants. The technical development therefore extended beyond map navigation to maintaining agreement between individual controls and shared outcomes.''',
zh='A1已完成昵称、三角色、确认与键盘移动；A2扩展为手机输入、二维码和共享状态。公共屏与手机分工不是A2临时改变；MMM Ball已在A1为浏览器参与提供参考。',
evidence='A1第14页原型流程→A2两手机/Firebase/公共大屏结构图；选一段真实共享状态代码。',
caption='Figure 2. The local A1 interaction flow was extended through phone input, QR access and synchronised player and activity data.',
source='Formal A1 PDF, pp. 6, 10, 14; project source; p5.js, Firebase Realtime Database and GitHub Pages documentation.')
# Preserve documented implementation problems; anchor their significance to the formal proposal.
D[2]['en']=D[2]['en'].replace('These failures affected the project’s central social condition.','These failures affected the distinction established in A1 between occupying the same space and responding to other participants.')
D[2]['source']='Formal A1 PDF, pp. 8, 10, 14; multiplayer implementation and recorded device feedback. No recruited-participant study is claimed.'
D[3]['en']=D[3]['en'].replace('The early phone layout allocated insufficient space to the map.','Moving beyond A1’s keyboard-based prototype required a usable portrait phone view. Early mobile layouts allocated insufficient space to the map.')
D[3]['source']='Formal A1 PDF, pp. 11–12, 14; phone-layout iterations and recorded device feedback.'
update(5,title='Developing the A1 Visual Direction',en='''A1 established a soft fantasy-game direction through pastel colours, chibi characters and illustrated 2D environments. The existing campus and character artwork provided the starting point for A2. New activity entrances needed to remain recognisable while fitting this established setting.

Several early generated options appeared overly saturated or detailed, leading to revisions towards softer colours, simpler forms and clearer outlines. Evaluation also considered each image within the map. The drink shop was enlarged, widened and repositioned; the fishing entrance used two rods and a lighter water surface without visible fish; the piano entrance was developed without a wooden sign.

These decisions developed the proposed visual direction through actual assets and placement. The campus functions as a setting for exploration rather than a fixed storyline. Familiar locations give activities context, while consistent proximity effects indicate opportunities to interact. The environment, responsive entrances and piano sound connect separate activities within one shared place, helping participants recognise both the setting and the actions available within it.''',
zh='明确柔和奇幻游戏风、粉彩与chibi角色已在A1第9页提出。A2写实际素材选择、删改、地图比例和接近效果，不声称从零提出统一画风。',
evidence='A1第9页moodboard局部＋真实图标前后版本＋最终地图位置检查。字体/色值若未逐项验证，不声称严格照搬A1色板和字体。',
caption='Figure 5. The A1 visual direction was developed through activity-image revisions, map placement and consistent interaction cues.',
source='Formal A1 PDF, pp. 9, 14; asset iterations and map-placement records. Original artwork and AI-assisted media require separate attribution.')
D[5]['source']='Formal A1 PDF, pp. 8, 10, for proposed shared activities; chat/wishing-tree implementation and development records. The message-board reference requires separate attribution if reproduced.'
D[6]['en']=D[6]['en'].replace('The drink activity distributes four decisions between two participants.','The drink activity develops A1’s proposal for short shared play by distributing four decisions between two participants.')
D[6]['source']='Formal A1 PDF, pp. 5, 8, 10; drink-maker implementation and feedback. The particular four-stage recipe was developed during A2.'
D[7]['source']='Formal A1 PDF, pp. 5, 8, 10; fishing implementation and recorded adjustments; GRDE2053 Session 2, pp. 8, 10.'
D[8]['source']='Formal A1 PDF, pp. 5, 10, for shared play; piano implementation and iPhone feedback; MDN Web Audio API documentation.'
D[9]['en']=D[9]['en'].replace('Development was informed by repeated device checks and feedback recorded during production.','A1 tested the local joining and navigation flow. A2 extended these checks to separate devices and shared activities, using repeated device checks and feedback recorded during production.')
D[9]['source']='Formal A1 PDF, p. 14; development/device feedback and available implementation records. Distinguish developer checks from independent audience research.'
update(11,title='Developing the Six Social UX Principles',en='''A1 identified Social Dynamics and Proactivity & Clarity as particularly important. A2 developed these intentions through implemented interactions and device feedback.

Sensory Feedback. Piano sound extends A1’s visual selection and movement feedback. The iPhone audio issue demonstrated the need to check the audible response independently.

Embodied Interaction. Phone-based movement and activity controls realise the proposed personal-controller role, although embodiment remains mediated through a touchscreen.

Social Dynamics. Alternating drink choices and divided fishing controls give both participants influence over one outcome. This develops A1’s distinction between co-presence and meaningful shared play, relating to the course’s discussion of computer-supported cooperative work.

Resilience & Motivation. Revised fishing settings and retry make the task more forgiving while retaining coordination.

Proactivity & Clarity. QR entry, proximity effects and turn notifications communicate where and how to participate, acting as signifiers.

Ethics & Inclusion. The non-human avatar option identified in A1 was retained. Sound control, exits and piano release further support choice and access. Phone ownership, connectivity, language, visual controls and public messages remain inclusion and privacy concerns.''',
zh='六原则在A1第15页已有预期应用；A2回顾它们怎样通过实现和测试发展。补回A1对非人类购物袋角色的包容性理由，但它是保留项，不说A2首次增加。名称按课堂原名，序号差异不当成概念变化。',
evidence='六格对比A1意图→A2实现；购物袋角色、轮次/靠近提示、分轴钓鱼、钢琴声音与让位可配小图。',
caption='Figure 11. A1’s proposed Social UX applications were developed through specific A2 interactions and evaluated against remaining barriers.',
source='Formal A1 PDF, p. 15; GRDE2053 Session 2, pp. 6, 8, 10; Session 3, p. 8; Session 5, slides 6–11.')
update(12,title='Evaluation Against the A1 Design Criteria',en='''A1 defined the experience as low-pressure, spontaneous, shared, playful and socially meaningful. The final prototype addresses these criteria through browser entry, optional activities and several ways to participate. Visitors can exchange messages, contribute a wish, coordinate a task or listen to another player. Each activity gives participants a potential reason to notice or respond to someone else.

The central evaluation question remains whether shared play leads to meaningful interaction between unfamiliar participants. Successful synchronisation or task completion cannot establish this alone. Observation would be needed to examine how players encounter one another, understand their roles and communicate during or after an activity. This follows A1’s social intention of creating an opportunity for connection without guaranteeing friendship.

A1 also proposed a public display with a nearby standing area and personal-phone control. Exhibition preparation must test this arrangement with the actual network and devices. The intended route of joining and completing one activity aims for approximately two minutes, including matching, but requires timed rehearsal. Phone activity screens must also be visible during demonstration, with a recorded fallback for connection failures.''',
zh='删除“文化提示在A2被弱化”这一错误比较。按正式A1第4页设计准则和第10页社会意图评估：能否给陌生人自然回应的理由，仍需参与者观察，不保证友谊。展场分工也已在A1第13页规划。',
evidence='A1设计准则—A2机制—仍待观察三列表；附A1第13页空间图与最终设备布置。',
caption='Figure 12. The final prototype is evaluated against A1’s design criteria, with social outcomes and exhibition conditions identified for further testing.',
source='Formal A1 PDF, pp. 4, 10, 13; final implementation/development record; A2 brief, pp. 2–4.')
# Preserve AI-use statement; new supplied A1 resolves only its own status, not image permissions or AI approval.
D[12]['source']='A2 brief, pp. 3, 5; project development and asset records. Formal A1 authorship is visible on its cover; original asset licences and AI approval remain separate matters.'
(O/'page-content.json').write_text(json.dumps(D,ensure_ascii=False,indent=2))
comparison=[['Aspect','Formal A1 evidence','A2 development'],['Social purpose','pp.2–4:low-pressure contact beyond familiar circles','Retained;implemented opportunities to notice,respond and cooperate'],['Concept stage','p.8:prompt-based screen already evolved into map/shared play','Implementation of that established direction,not a new conceptual shift'],['Technical prototype','p.14:single-user p5.js,nickname,three avatars,confirmation,arrow-key movement','QR access,phone control,independent players,shared display and activity state'],['Shared activities','pp.8,10:short activities proposed,not implemented','Specific chat,wishes,drink,fishing and piano interactions developed'],['Visual direction','p.9:pastel fantasy-game/chibi/2D setting','Actual entrance images,placement,scale and proximity cues refined'],['Social UX','p.15:proposed applications,especially Social Dynamics/Clarity','Mechanisms implemented and reviewed through device feedback'],['References','pp.5–6:21 Swings and MMM Ball','Cooperation and browser-entry relevance retained;Impulse not a formal A1 case'],['Tokens','Not included in this formal15-page submission','Do not describe token removal as a formal A1-to-A2 change']]
with (O/'A1-to-A2-comparison.csv').open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(comparison)
for bilingual,name in [(False,'Development-content-EN.docx'),(True,'Development-content-EN-ZH.docx')]:
 doc=Document();doc.sections[0].top_margin=doc.sections[0].bottom_margin=Inches(.7);doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10 if bilingual else 11)
 for i,p in enumerate(D):
  if i:doc.add_page_break()
  doc.add_heading(f"{p['page']:02d}  {p['title']}",0)
  for para in p['en'].split('\n\n'):doc.add_paragraph(para)
  if i in [0,1,4,10,11]:
   sources={0:'A1 source: Zhao, Shared Campus Map, pp. 2–4, 8, 10, 14.',1:'A1 source: pp. 6, 10, 14.',4:'A1 source: pp. 9, 14.',10:'A1 source: p. 15. Course sources: Session 2, pp. 6, 8, 10; Session 3, p. 8; Session 5, slides 6–11.',11:'A1 source: pp. 4, 10, 13.'}
   doc.add_paragraph(sources[i],style='Caption')
  if bilingual:
   doc.add_heading('中文说明与证据安排（不放入英文正文）',2);doc.add_paragraph(p['zh']);doc.add_paragraph('配图：'+p['evidence']);doc.add_paragraph(p['caption']);doc.add_paragraph('来源依据：'+p['source'])
 doc.save(O/name)
md='# Shared Campus Map — Development Documentation\n\n'
for p in D:md+=f"## {p['page']:02d}. {p['title']}\n\n{p['en']}\n\n"
(O/'Development-content-EN.md').write_text(md)
print('Updated13pages;',sum(len(re.findall(r"\b[\w’–-]+\b",p['en'])) for p in D),'English words.')
