from pathlib import Path
import json,re,csv
from docx import Document
from docx.shared import Pt
P=Path('/workspace/presentation-reviewed');m=json.loads((P/'presentation-content.json').read_text());words=sum(len(re.findall(r"\b[\w’–-]+\b",x['en'])) for x in m)
d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(11)
d.add_heading('Shared Campus — 展示准备与复核',0)
d.add_paragraph('本说明只用于准备，不朗读、不投影。正文讲稿、PPT备注和时间表由同一份presentation-content.json生成。')
d.add_heading('1. 时间：6分45秒为排练目标，不是实测结果',1)
d.add_paragraph(f'英文正文{words}词，75秒真实展示视频；现场环节55秒，包含约14秒引入与约41秒操作。第7页80秒包含视频引入。视频和现场操作都计入5–7分钟。切勿把中文解释再读一遍。')
d.add_paragraph('逐页时间是预算。100词/分钟时正文约296秒，加75秒视频与约41秒现场非朗读操作，约6分52秒，尚未另算翻页停顿；110词/分钟时约6分25秒，120词/分钟时约6分03秒。因此语速约110词/分钟再完整排练，比仅看6:45标签可靠。不能保证每个人照读都在7分钟内。')
t=d.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['页','内容','预算']):c.text=v
for x in m:
 for c,v in zip(t.add_row().cells,[str(x['slide']),x['title'],x['time']]):c.text=v

d.add_heading('2. 75秒真实使用视频：连续分镜',1)
shots=[('0–10秒','公共展览情境；大屏、两位参与者与手机同框。'),('10–20秒','扫码、昵称/角色、进入并看到彼此。'),('20–35秒','文字气泡交流；饮品轮流决定，双方看到结果。'),('35–50秒','钓鱼不同方向控制、浮标涟漪与共同结果。'),('50–60秒','挂愿望；另一人点击查看并祝福。'),('60–70秒','一人弹钢琴、另一人开启声音聆听；保留琴声。'),('70–75秒','回到共享校园全景。')]
for a,b in shots:d.add_paragraph(a+'：'+b)
d.add_paragraph('当前还没有真实视频；第7页是占位。使用真实参与者、获得适当拍摄同意。控制测试截图只能说明界面，不能代替真实观众研究。实拍素材决定最终分镜，不虚构互动或结果。')

d.add_heading('3. 现场演示：先解决观众看不看得清',1)
d.add_paragraph('饮品制作面板在手机上，公共地图大屏并不等于手机画面。提前安排搭档操作第二台手机。优先用实物展台/摄像头拍两台手机；或投屏一台手机并能切换展示另一台。把手机屏幕呈现给老师，不能只让参与者低头操作。')
d.add_paragraph(m[7]['cue'])
d.add_paragraph('公共地图固定入口：https://itsangelia.github.io/public-map-project/campus-f3cd85c8e3d0eaa9.html?mode=display&room=showcase\n手机入口：https://itsangelia.github.io/public-map-project/campus-f3cd85c8e3d0eaa9.html?room=showcase\njoin-qr.png指向手机入口。演示前在展场网络、实际投屏设备上排练。链接与既有正式版本一致，本轮未在线验证可达性。')
d.add_paragraph('预先准备备用录像。若连接失败，不持续在台上调试；说明故障并切换录像，明确它是录像，不冒充live。两分钟参与路径属于设计目标，扫码/匹配等待也需计时实测。')

d.add_heading('4. 老师要求与页码',1)
rows=[('社会焦点与人本议题','2：国际学生初次接触；Connection为主，Joy/Reflection为辅。'),('共享参与与有意义的选择','3–6：地图/聊天/许愿祝福；饮品共同决策；钓鱼必要合作；钢琴演奏与聆听。'),('叙事、氛围或体验设计','10：校园提供场景，视觉和声音连成一个体验。'),('六原则及课堂理论','9：六个准确名称，CSCW接分工合作，signifiers接发光与轮次提示。'),('开发、迭代及重要变化','11：与初始概念比较、多人可见性、手机视角与钓鱼反馈。'),('1–2分钟视频与现场演示','7：75秒视频；8：55秒现场环节。视频未提供，需实拍并插入。'),('引用与AI透明','13–14、来源登记、备注；批准/原作者信息仍待核实。')]
for a,b in rows:d.add_paragraph(a+'：'+b)
d.add_paragraph('课件出处：六原则见Session 2 PDF p6、Session 3 PDF p8、Session 5 slides6–7。CSCW见Session2 p8；signifiers见Session5 slides9–11。仅据课堂课件讨论理论，不假称阅读全部原著。')
d.add_paragraph('附带文件Social Interactive.pptx支持初始概念比较，但是否为正式A1提交版尚未确认。当前说initial concept，不说已核实的正式A1提交版本。')

d.add_heading('5. 提交前仍待完成',1)
for x in ['封面姓名、实际75秒视频。','在展场完成一次带搭档/投屏的完整排练，实测5–7分钟；单活动参与路径实测约两分钟。','原地图/角色的作者、来源与许可；具体AI素材与工具记录；按unit outline规定格式统一引用。','核实导师允许的AI范围及批准，提交真实AI使用声明；PPT里的说明不能替代声明。','独立development documentation PDF最多15页，不计标题和参考文献；本PPT不替代。','Curtin coversheet、要求的完整录制MP4/SharePoint分享链接与导师访问权限，以及Week10里程碑记录。']:d.add_paragraph('□ '+x)
d.add_paragraph('Brief和rubric中的5–7分钟、视频/现场与过程文档是教师要求。不会把素材授权或导师批准的未知状态当成已经完成，也不把自动化检查当参与者研究。')
d.save(P/'Teacher-requirements-and-rehearsal-ZH.docx')
rows=[['asset_or_media','use','known_source','action_before_submission'],['Prototype screenshots','PPT1,5,6,8,10','Project controlled browser captures,2026','Not participant-study evidence'],['map draft.png / avatars','Shared world and characters','Supplied for the project; original provenance not established','Confirm original creator,source and licence'],['Activity icons / fish media','Activity entrances and catch results','AI-assisted generation/refinement; some supplied originals','List each asset,original source,tool and adaptation'],['Piano audio','Map performance','Project Web Audio implementation','Disclose code assistance; do not claim recorded instrument samples'],['Six Social UX principles','PPT9','Session2 p6;Session3 p8;Session5 slides6–7','Use unit referencing style'],['CSCW / signifiers','PPT9','Session2 p8;Session5 slides9–11','Cite lecture discussion; primary sources not assumed read'],['Initial concept','PPT11','Social Interactive.pptx slides5,7,8','Confirm relation to actual A1 submission'],['Showcase footage','PPT7','Not supplied','Record real use,obtain appropriate consent,insert75-second film']]
with (P/'Media-and-AI-source-register.csv').open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
(P/'READ-ME-FIRST.md').write_text(f'''# Shared Campus — reviewed presentation\n\nEditable15-slide PPT and a matching PDF. Slides1–14 are the talk;15 is an untimed submission appendix. English speech with matching Chinese explanations is in DOCX/Markdown and PowerPoint notes.\n\nTiming target6:45;{words}spoken English words,75-second video,55-second live segment including spoken introduction. Budgets are estimates: rehearse on the actual devices. The Markdown, DOCX, notes and timing chart share one JSON source.\n\nStill required: presenter name,real showcase footage,actual demo projection setup and rehearsal,source/licence confirmation,unit citation style,tutor AI approval and declaration. Initial concept deck is supplied;its status as the formal A1 submission is unconfirmed. No fabricated participant study or proven belonging outcome.\n\nNo application code or deployed game behavior changed. The package does not replace the separate process PDF,coversheet or MP4/SharePoint submission required by the brief.\n''')
(P/'Review-findings-ZH.md').write_text('''# 第二轮审核结果\n\n已修正：\n- Markdown开头60秒与正文55秒冲突；统一55秒。\n- 视频许愿树片段误写50–55秒；恢复50–60秒，分镜0–75秒连续。\n- 技术页与结束页语速突快；重分时间为6:45预算，不冒称实测。\n- 饮品面板在手机而公共地图不显示；备注明确搭档与投屏/摄像头方案。\n- 钓鱼不是拖动浮标；用方向按钮与有限触屏具身互动的准确描述。\n- CSCW和signifiers不再孤立报名称，分别解释必要分工及交互提示。\n- 昵称与退出不代表已实现全面隐私/无障碍；明确声音控制、选择及公平使用的具体贡献与剩余障碍。\n- 两分钟改为短路径设计目标，不当成已有实测结果。\n- 原概念的文化交流意图在成品中较间接；加入这一实质局限，避免只肯定功能。\n- 观众可见页的作者待办移到准备说明/备注；未核实素材仍标明待核实状态。\n- 素材登记页码更新到1/5/6/8/10，7为视频占位。\n\n尚未完成：真实视频、展场投屏和完整排练、姓名、完整素材原作者/许可、导师AI批准核实。讲稿修订完成并不等于整份作业已可提交。\n''')
print('Support rebuilt from current metadata.')
