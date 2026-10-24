from pathlib import Path
import json
from docx import Document
from docx.shared import Pt, Inches
P=Path('/workspace/development-documentation')
D=[]
def page(title,en,zh,evidence,caption,source):D.append(dict(page=len(D)+1,title=title,en=en,zh=zh,evidence=evidence,caption=caption,source=source))
page('From Initial Concept to Working Experience',
'''Shared Campus is a browser-based social experience that connects personal phones to a shared illustrated campus on a public display. My initial concept focused on international students and the awkwardness of approaching someone unfamiliar. QR entry, avatars, free exploration and cooperative activities were already part of that concept; they were not new ideas introduced during final development.

The initial proposal also described everyday cultural prompts, shared visual outcomes and a possible token system. The final version retains the connection goal, but develops it through chat, a wishing tree, collaborative drink-making, cooperative fishing and a shared piano. Tokens were not implemented. Participation can be immediate, such as chatting or playing together, or asynchronous, such as reading and blessing a wish.

The main development challenge was turning a proposed shared space into a usable experience across separate devices. This document records the implementation, problems, revisions and remaining limitations. It does not treat a completed prototype as evidence that social belonging has improved.''',
'''这里用很短的篇幅接回初始概念：国际学生的初次连接、二维码、角色、自由探索与合作，本来就已经存在。A2真正完成的是把它们落实成能运行的联机体验。文化提示与代币没有按原样实现，应如实说明。不要把这页写成再次提出项目。''',
'左侧放Social Interactive概念稿第5或7页局部；右侧放最终地图。用短标注区分保留、发展、未实现。',
'Figure 1. Initial concept and final prototype: the shared-map approach was retained, while specific activities developed during production.',
'Initial concept: Social Interactive.pptx, slides 2, 5, 7–8; project development record.')
page('Development Stages and Technical Structure',
'''Development progressed from the existing p5.js map to shared player presence, communication and individual activity systems. The illustrated map and supplied characters remained the visual foundation. I added activities incrementally and returned to the shared-map experience when new problems appeared.

The system separates the shared world from each device’s view. Phones provide individual input and follow their own characters; the public display presents the common map. Firebase Realtime Database distributes player and activity data, while GitHub Pages provides a reachable web entry point. This makes it possible for separate phones to participate without installing an application.

The architecture needed to support more than movement. Messages, wishes, drink turns, fishing controls and piano occupancy each required shared state. A coherent room is therefore central to the experience: participants need to join the same room and receive compatible state. Successful local rendering alone cannot establish that multiplayer interaction works.''',
'''把过程分成已有地图、多人连接、交流功能、合作活动、整体完善几个阶段。技术图解释谁负责什么，不需要大段代码。特别强调手机显示自己的视角，不代表每台手机在玩互不相关的地图。''',
'开发阶段时间线，不填未知日期；结构图：两台手机↔Firebase↔公共大屏。可附一小段真实玩家数据结构，隐藏任何身份标识。',
'Figure 2. Personal controls and a public overview connect through shared room data; device-specific cameras do not replace shared world coordinates.',
'Project source: sketch.js and multiplayer/activity modules; p5.js, Firebase and GitHub Pages documentation.')
page('Making Other Players Reliably Visible',
'''A recurring problem was that a second device could enter the map without seeing the other player, or the display appeared to contain only one participant. This was especially disruptive because seeing another person is the starting condition for the social experience. A connection label did not, by itself, demonstrate that all players had been received and rendered.

The multiplayer implementation uses distinct player identities and connection records rather than treating every phone as one shared character. Reviewing the join flow, room state and player updates became a priority. The useful check was an independent display and two phones, each controlling a separate character, rather than several views of one session.

Subsequent device feedback confirmed occasions when both players became visible and could interact. Earlier failures also returned during development, so these confirmations should not be presented as proof of reliability under every network condition. The important lesson was to verify shared presence again when introducing a new feature, rather than assuming that a working local screen guaranteed a working shared experience.''',
'''写真实反复出现的第二台设备问题。不要编造一个没有证实的唯一根因。可以解释当前独立身份与连接结构，但不能说仅某一次改动就永久解决所有情况。把连接状态与真正看见其他玩家区分开。''',
'若保留了早期“players: 1”与后期双玩家画面，做前后对比。没有旧图就用真实记录摘录，明确是开发反馈；新增复测要注明新测试日期。',
'Figure 3. Shared presence was checked through independent player control and visibility on another device, rather than through a connection label alone.',
'Development conversation and multiplayer.js. Reported device checks are not a formal participant study.')
page('Refining the Phone View and Movement',
'''The early phone layout left a large empty area below the map. The character appeared small, and the map did not extend far enough behind the virtual controls. Although the world was visible, the screen did not give enough space to the part of the experience the player was actively using.

The revised layout fills the available map area and places the joystick over the scene. The phone view was enlarged by 35%, and movement speed was reduced by 20% after visual trials. These changes made the local view more readable while reducing the feeling of moving too quickly through the enlarged scene. The public display retained its overview role.

This iteration showed that the same world needs different presentations on a phone and a large screen. It also exposed a trade-off: increasing local scale improves character visibility but reduces how much surrounding space can be seen. The final values reflect development feedback and visual preference; they are not the result of a measured accessibility or navigation study.''',
'''这是非常明确的前后变化：空白太多、人物太小→铺满地图、35%视角放大、20%减速、摇杆叠在地图上。解释视觉放大与视野范围之间的取舍，不直接声称已经通过可用性研究。''',
'早期手机截图与最终手机截图并排，同尺寸展示；圈出地图边界、角色大小和摇杆位置。',
'Figure 4. The phone layout changed from an undersized map with unused space to a larger local view with controls over the scene.',
'Development feedback and phone-interface screenshots; sketch.js and styles.')
page('Developing a Consistent Visual Language',
'''The existing illustrated campus and characters established the visual direction. New activity entrances needed to belong to that world while remaining recognisable as interactive objects. I reviewed multiple versions rather than accepting the first generated image as a finished asset.

Feedback repeatedly identified excessive saturation, unnecessary detail and a strong “AI-generated” appearance. Revisions moved towards softer colours, simpler forms and clear outlines. Individual decisions also mattered: the drink shop was widened and repositioned, the fishing entrance used two rods and a lighter water surface without visible fish, and the piano entrance avoided a wooden sign.

Scale and placement were refined within the map, rather than judging the icons only as isolated illustrations. Proximity effects provided a consistent invitation to interact. This process combined asset selection, visual adjustment and contextual checking. The original map and avatar artwork were retained; their authorship should remain separate from the AI-assisted activity assets.''',
'''展示你怎样选风格、拒绝不合适版本，才是个人设计贡献。不要说地图和角色都是自己画的，也不要把所有图都写成同一来源。不同入口的小决定可以作为具体证据。''',
'选2–3组真实早期/最终图标，不必把所有生成图都放上。附一张最终地图中的位置与比例检查。',
'Figure 5. Selected revisions reduced visual complexity and adjusted scale so activity entrances could fit the established campus style.',
'Development asset iterations. Attribute supplied artwork and AI-assisted media individually in the media register.')
page('Chat and Wishes: Two Speeds of Participation',
'''Chat and the wishing tree support different forms of contribution. Typed chat appears above a character and disappears after five seconds, providing an immediate way to address other players without keeping a permanent speech bubble over the map.

The wishing tree retains a contribution for later discovery. Players open the tree interface, write a wish and hang a note. Other visitors can view a short preview, open the full message with its nickname and date, and add a blessing. The preview and gentle movement developed from a specific concern: visitors needed a reason to click someone else’s wish. Showing the entire message on the tree removed that reason, so the preview was shortened to roughly four or five English words.

The blessing button adds a small response that does not require composing another message. These features create opportunities for expression and acknowledgement, but their availability does not demonstrate that strangers feel connected. Publicly readable text also creates privacy and moderation responsibilities.''',
'''这页重点是为什么把即时聊天与持续纸条分开，以及预览为什么要短。祝福提供低成本回应，但不能把按钮次数自动当成情感连接的证据。公开文字的风险留到最终反思一起讨论。''',
'三个局部：5秒气泡、树上短预览、展开后的愿望与祝福。用箭头说明从吸引点击到查看完整内容。',
'Figure 6. Immediate chat and persistent wishes offer different rhythms of participation; short previews invite discovery without revealing the full note.',
'chat.js, wishing-tree modules and development feedback on previews and blessings.')
page('Drink-Making: Sharing Decisions',
'''The drink shop turns collaboration into a sequence of small decisions. Two players match with visible nicknames and character portraits. They alternate across four stages: base, flavour, topping and decoration. Each stage has three options, allowing 81 recipe combinations. The final named drink reflects choices made by both participants.

Early visual feedback identified that changing an option did not produce a sufficiently noticeable change in the drink. The development priority was therefore to make ingredient changes distinguishable while keeping the approved visual style. A “Your turn” notification clarifies who can act. The selection cards were kept upright, while unnecessary cup logos, straws and base shadows were removed.

The result screen enlarges the completed drink and provides a clear return to the map. There is no individual winner: each player influences the same outcome. This is a concrete mechanism for shared participation, although conversation between the players remains optional rather than guaranteed.''',
'''讲清双方怎样共同影响结果，再写“点击后变化不明显”这一真实反馈和修订。风格统一、轮次清楚、最终共同成品都是设计决定，不只是游戏规则。''',
'匹配→轮到我→四步选择的局部对比→放大的结果。选两种差异明显的真实成品来证明反馈可见。',
'Figure 7. Alternating choices give both players a contribution to one shared drink; clear visual changes communicate the effect of each decision.',
'Drink-maker implementation and recorded interface feedback.')
page('Fishing: Balancing Cooperation and Forgiveness',
'''Fishing requires two players to coordinate one float. One controls up and down, while the other controls left and right. Neither controls both axes alone. After a countdown, they aim the float at moving ripples to build shared progress. The target is represented by water movement rather than a visible fish, allowing the catch to remain a surprise.

Development feedback identified jerky movement and an overly difficult challenge. Revisions included smoother visual movement and more forgiving settings: the alignment range increased by 10%, target changes were spaced at 3.5 seconds, and progress loss when misaligned was reduced to 4% per second. Progress begins at 20% and increases by 10% per second while aligned; reaching zero fails the attempt.

Success reveals one of four named fish, while failure shows a static escape result and offers another attempt. These changes support continued play without removing the need for cooperation. Reducing difficulty addresses challenge, but it should not be used as a substitute for fixing network or rendering problems.''',
'''钓鱼是迭代证据最强的一页。把卡顿与难度分开：平滑处理针对移动感，容错和进度下降调整针对挑战。不能写成靠降难度修复网卡。鱼图边缘处理和命名可作为小型补图。''',
'过程截图包括分工控制、半浸浮标、涟漪、进度及成功/失败。小表列设置修改前后；未知的旧值不填。',
'Figure 8. Separate input roles make coordination necessary, while revised timing and tolerance reduce avoidable frustration.',
'Fishing modules and requested difficulty adjustments; Session 2, pp. 8, 10, for cooperation and forgiving interaction.')
page('Piano: Performance, Listening and Fair Access',
'''The shared piano introduces participation through both performance and listening. One person occupies the instrument and plays a three-row, 21-key interface. Other players can listen after enabling sound, while an occupied notice identifies the performer. The activity icon responds visually during use.

Technical development included database permission setup and an iPhone audio issue: keys changed colour, but no sound was heard. This distinguished a working visual response from successful audio output. After changes to audio activation and recovery, device feedback confirmed that sound could be heard. This confirmation applies to those checks, not every possible browser or device.

The interaction was also changed from a separate page to a panel within the map, keeping the shared setting visible. Controls were moved away from the exit to reduce accidental taps. Eight seconds without a note releases the seat, preventing an idle session from indefinitely holding the instrument. Sound control, clear exits and automatic release support participant choice and shared access.''',
'''钢琴有两类开发：技术出声和权限，以及参与方式/公平使用。按键变色不等于音频正常。退出误触、地图内面板、只显示演奏者通知和8秒让位，都能连接到明确理由。''',
'权限或无声问题记录、最终地图内键盘、占用昵称提示、演奏中图标。显示实际截图，不重现含私人信息的控制台。',
'Figure 9. The piano developed into an on-map activity with audible feedback, clear occupancy and an inactivity release.',
'Shared-piano modules and iPhone device feedback; MDN Web Audio API documentation.')
page('Testing, Feedback and Responsive Changes',
'''Testing used independent browser/device views and repeated development feedback. Checks included whether two phones could control different characters, whether another player could receive messages, whether activity outcomes agreed, and whether iPhone piano audio worked. These checks address functionality and interaction clarity; they are not a formal study of social impact.

The most useful feedback identified a concrete mismatch between the intended and actual experience. Examples included absent players, excessive mobile whitespace, weak ingredient differences, jerky fishing and silent piano keys. Each issue led to a specific revision that could be checked again. Features were not assumed to remain reliable simply because an earlier version had worked.

The evidence should show the sequence of issue, change and re-check. User confirmations in the development conversation support reported improvements, but do not establish audience numbers, universal compatibility or long-term reliability. A final exhibition rehearsal still needs the actual network, display, two phones and a visible demonstration setup. Waiting for a partner must also be included when checking the intended short encounter.''',
'''这一页用真实表格集中回顾，不再重复功能介绍。证据分成开发对话、设备检查和程序检查；不把它们包装成独立用户研究。没有导师评价或问卷就不编。''',
'用下方附表中的记录，配2–3个真实截图。复测未记录的项目写“需补充”，不要写pass。',
'Figure 10. Development evidence connects specific problems to revisions and reported re-checks; remaining exhibition checks are recorded separately.',
'Development conversation, repository test records where available, and actual device screenshots.')
page('Reflecting on the Six Principles for Social UX',
'''The course framework provides a way to evaluate both the strengths and limits of the final experience.

Sensory Feedback: piano sound responds to input, supported by visible key responses. Embodied Interaction: touch controls connect physical input to movement, although embodiment is limited to a phone rather than whole-body interaction. Social Dynamics: alternating drink choices and split fishing controls give participants influence over a shared outcome. This connects to the course’s discussion of computer-supported cooperative work.

Resilience and Motivation: easier fishing, clear results and retry support continued participation. Proactivity and Clarity: proximity glow and turn prompts act as signifiers, communicating available actions and current responsibility. Ethics and Inclusion: sound control, exits and piano release support choice and fair access. However, dependence on phones, internet, language and visual controls creates barriers, and public messages require care.

These examples do not mean all six principles are fully resolved. In particular, touchscreen input is a limited form of embodiment, and the prototype has not established comprehensive accessibility.''',
'''名称必须与课件一致。每项都对应一个设计和理由，不写“已经完全满足”。CSCW接合作分工，signifiers接提示，不单独报理论名词。具身与包容性的不足需要明确。''',
'做六格，每格放原则名、项目局部、理由及一句局限；不要只用无项目证据的装饰图标。',
'Figure 11. The six Social UX principles are used to reflect on concrete interactions and unresolved participation barriers.',
'GRDE2053 Session 2, pp. 6, 8, 10; Session 3, p. 8; Session 5, slides 6–11. Theory discussed through the supplied lectures.')
page('Final Evaluation and Exhibition Readiness',
'''The final prototype provides several levels of social involvement: noticing other avatars, exchanging a short message, leaving a wish, coordinating a task or listening to music. Participants can choose an activity instead of following one compulsory sequence. A single short encounter, rather than completion of the whole map, is the intended exhibition unit.

There is also a significant limitation in the concept’s evolution. The initial prompts explicitly invited everyday cultural sharing. The final activities offer a reason for first contact but address cultural exchange less directly. It would therefore be premature to claim that the project improves cross-cultural belonging. A useful next evaluation would observe whether unfamiliar participants begin together, understand their roles and choose to communicate beyond the task.

Exhibition preparation requires a reachable QR link, stable network, charged phones, a suitable public display and a visible way to demonstrate phone activity screens. The approximately two-minute participation target still needs timed rehearsal, including joining and matching. A recorded fallback should be available without being presented as live interaction.''',
'''评价不能只说“项目很好”。承认原先的文化分享变得更间接，再提出具体可观察的问题。约两分钟是待实测目标。演示饮品需要投屏或摄像头，公共地图不能替代手机制作画面。''',
'最终展览布局图：公共地图、二维码、两台手机、投屏/摄像头方案。旁边列已确认条件与仍待确认事项。',
'Figure 12. Exhibition planning separates the shared public view from individual controls and treats participation time as a target to validate.',
'Initial concept and development record; A2 brief, pp. 2–4.')
page('AI Use and My Design Contribution',
'''Generative AI assisted parts of the prototype code, debugging, iteration planning, selected activity artwork and preparation of written material. These uses need to be documented as part of the project’s authorship rather than described as entirely independent coding or hand-drawn production.

My contribution included setting the social direction, defining activity rules, selecting and rejecting visual options, specifying revisions and reviewing the experience through device feedback. Examples include rejecting overly saturated icons, requesting more visible ingredient changes, refining fishing difficulty and keeping the piano interaction within the map. AI outputs were adapted within the existing project rather than treated as finished decisions.

The supplied map and character artwork must be credited separately from generated or refined activity assets. Individual media records should identify the source, tool, use and subsequent changes. The original creators and licences of supplied assets still require confirmation. This draft does not assert that tutor approval has been obtained: the submission must follow the unit’s permitted-use rules and include the required AI declaration.''',
'''如实说明AI用途，同时用具体选择与修改体现你的贡献。不要写成所有代码都是独立编写，也不要把AI说成只润色文字。导师批准、图像来源未知就标为待核实；提交前按真实情况修改最后一段。''',
'建议一张“AI输出→你的反馈/决定→最终整合”的真实例子，加一张素材/工具登记表。工具具体型号与日期未知则不编。',
'Figure 13. Authorship is documented through the relationship between assisted outputs and the designer’s selection, revision and integration decisions.',
'A2 brief, pp. 3, 5; project development and asset records.')
(P/'page-content.json').write_text(json.dumps(D,ensure_ascii=False,indent=2))
# Clean English copy: evidence instructions are in the bilingual companion, not in page body.
for bilingual,filename in [(False,'Development-content-EN.docx'),(True,'Development-content-EN-ZH.docx')]:
 doc=Document();sec=doc.sections[0];sec.header.paragraphs[0].text='Shared Campus | Development Documentation — Content Draft'
 sec.top_margin=sec.bottom_margin=Inches(.7);doc.styles['Normal'].font.name='Arial';doc.styles['Normal'].font.size=Pt(10 if bilingual else 11)
 for i,x in enumerate(D):
  if i:doc.add_page_break()
  doc.add_heading(f"{x['page']:02d}  {x['title']}",0)
  for para in x['en'].split('\n\n'):doc.add_paragraph(para)
  doc.add_paragraph('Source basis: '+x['source'],style='Caption')
  if bilingual:
   doc.add_heading('中文理解与写作说明（不放入英文成品）',2);doc.add_paragraph(x['zh']);doc.add_paragraph('建议证据：'+x['evidence']);doc.add_paragraph('图片说明：'+x['caption'])
 doc.save(P/filename)
md='# Shared Campus — Development Documentation\n\nEnglish body copy for 13 planned content pages. Cover and references are separate. Layout and evidence must be added before PDF submission.\n\n'
for x in D:md+=f"## {x['page']:02d}. {x['title']}\n\n{x['en']}\n\n*Source basis: {x['source']}*\n\n"
(P/'Development-content-EN.md').write_text(md)
(P/'中文理解与配图说明.md').write_text('# 正文对应与配图说明\n\n'+''.join(f"## 第{x['page']}页 {x['title']}\n\n{x['zh']}\n\n建议证据：{x['evidence']}\n\n英文图注：{x['caption']}\n\n" for x in D))
print('Pages:',len(D),'English words:',sum(len(x['en'].split()) for x in D))
