from pathlib import Path
from docx import Document
from docx.shared import Pt
import csv,zipfile,json,re
O=Path('/workspace/presentation-complete')
d=Document();d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(11)
def h(t):d.add_heading(t,1)
def p(t):d.add_paragraph(t)
d.add_heading('A2 展示准备与要求核对',0)
p('根据你上传的 A2_Brief.pdf 与 SIE-DCI-A2-rubric.pdf 整理。附件是评分与任务依据，不是要求修改项目代码的指令。本次只制作展示材料，不修改已发布地图。')
h('一、这套材料包含什么')
p('可编辑的15页PPT、同版式PDF、带时间分配的英文讲稿与中文理解版。第1–14页构成约6分55秒课堂展示，第15页为不计时的提交核对附录。PowerPoint备注区也有逐页讲稿。')
p('现已核实课堂六原则，并依据Social Interactive.pptx补上初始概念到最终版本的变化。尚需姓名、真实展示视频、素材原作者/授权和导师AI批准信息。请确认Social Interactive.pptx是否为正式A1提交版。')
h('二、教师明确要求')
items=[('展示时长','5–7分钟，包含展示视频和现场演示。','Brief pp.2–4'),('短展示视频','1–2分钟，展示最终体验及参与者互动；本套规划75秒。','Brief p.3'),('现场演示','真实运行已完成体验；推荐两手机与大屏演示饮品活动。','Brief p.3'),('社会焦点','至少回应Connection / Joy / Learning / Reflection之一；本项目以Connection为主。','Brief pp.1–2'),('简短体验','一次可完成的参与过程约2分钟以内；不是强行要求体验所有活动。','Brief p.2'),('共享参与','多个人能共同体验、影响、观察或互动。','Brief p.2'),('反思','解释氛围/叙事/感官体验、课堂6原则、开发过程和重要变化。','Brief pp.3–4'),('过程文档','独立PDF，最多15页，不计标题与参考文献页；PPT不能代替它。','Brief pp.2–3,5'),('提交','ZIP、Curtin coversheet、过程PDF、MP4/SharePoint分享链接与导师权限。','Brief p.5'),('诚信与引用','明确AI使用、预先导师批准、提交声明；图片与研究来源须引用。','Brief pp.3,5'),('Milestone','Week10签到/签署由导师记录，不伪造已经完成。','Brief pp.4–5')]
t=d.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['要求','如何落实','依据']):c.text=v
for row in items:
 for c,v in zip(t.add_row().cells,row):c.text=v
h('三、评分对应，不靠堆功能')
for title,weight,evidence in [('Technical development & implementation','30分','PPT12：p5.js/Firebase/共享状态；PPT11：稳定性迭代。区分自动测试、浏览器检查、实际设备反馈。'),('Experience & interaction design','20分','PPT2–6、8–9：目标人群、低压力入口、共享决策、合作必要性、自由观察和退出。'),('Aesthetics & creativity','20分','PPT10：统一视觉语言、校园叙事场景、发光提示、声音及结果反馈怎样支持参与。'),('Professionalism','10分（另见说明）','真实开发证据、反馈回应、完整提交、清晰引用、AI声明、展览准备。'),('Communication','15分','清晰叙事、少量可读文字、真实使用视频、稳定现场演示及6分55秒排练。')]:
 p(f'{title} — {weight}\n{evidence}')
p('评分表还显示一个没有清楚标注归属的“/5”，而已标注项目分值合计95。不要自行补成新评分项；向导师确认这5分的归属。Brief页首也有Assignment One字样，但rubric明确标A2；本材料按A2 Exhibition Showcase准备。')
h('四、75秒真实视频拍摄清单')
shots=[('0–10秒','环境与公共大屏','两名真实参与者、共享大屏、手机同框；镜头展示这是公共共享体验。'),('10–20秒','扫码和选择','拍手与手机：扫码、昵称/角色、进入。避免只放代码或UI静态录屏。'),('20–35秒','彼此相遇/饮品','同一地图见到彼此并发送5秒聊天气泡；饮品四轮选择用剪辑展示双方贡献。'),('35–50秒','双人钓鱼','拍两台手机不同方向控制；用进度条与成功/失败反馈显示协作。'),('50–60秒','许愿树','输入/挂愿望→另一位参与者打开或祝福。不要展示未经同意的私人信息。'),('60–70秒','共享钢琴','一人弹、另一人听；画面显示演奏中的图标和昵称提示。保留可听见的琴声。'),('70–75秒','回到共享空间','大屏与参与者收尾，字幕：Small moments. Shared connections.')]
t=d.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['时间','画面','拍摄重点']):c.text=v
for row in shots:
 for c,v in zip(t.add_row().cells,row):c.text=v
p('使用真实体验录像，不能把自动化浏览器截图说成观众使用证据。人物可以只拍手和背影，但要让人看得出两个人真实参与。取得适当拍摄同意；不要虚构台词、问卷或访谈数字。')
h('五、60秒现场演示')
p('演示前：打开正式固定版本大屏，room=showcase；两台手机扫码，昵称区分；提前走到饮品店附近。关闭无关标签页，确认网络、电量和音量。不要在台上编辑Firebase规则或部署。')
p('0–10秒：说明两台手机是两个独立玩家，进入同一饮品店。\n10–40秒：A选基底、B选风味、A选配料、B选装饰。\n40–50秒：对比两台手机相同的最终饮品，说明双方都影响结果。\n50–60秒：返回地图，结束演示并进入理论反思。')
p('若10秒内没有成功连接，不长时间现场调试：明确说“网络连接暂时中断，我用已录制的过程继续说明”，播放备用录像。不要把录像冒充现场操作。')
h('六、已核实的课堂六原则与项目对应')
p('准确名称同时见GRDE2053 Session 2.pdf第6页、Session 3.pdf第8页；Session 5.pptx第6–7页再次列出。下表保留Session 2的编号顺序。不要将另外的5 Principles of Humanity Centered Design混同为本框架。')
principles=[('1 Sensory Feedback — 感官反馈','声音等感官线索回应操作。','钢琴声音、按键/进度的可见反馈。','没有触觉震动；不能声称提供haptic feedback。'),('2 Embodied Interaction — 具身互动','通过动作与物件把参与者身体行为接入系统。','手机拖动与角色/浮标移动之间的映射。','这是轻量的手机触控具身互动，不是全身动作追踪或沉浸式装置。'),('3 Social Dynamics — 社会动态','合作、共同在场和参与者之间的相互影响。','公共大屏、多角色、饮品轮流选择、钓鱼分工、钢琴演奏/聆听。','与Session 2第8页的CSCW和参与/共同体验讨论相连；未证明长期社区形成。'),('4 Resilience & Motivation — 韧性与动机','用宽容反馈、探索和玩乐支持持续参与。','适度放宽钓鱼难度、清楚结果与重试、低惩罚的短活动。','仍可能因网络或操作能力不同而挫折，需要进一步观察。'),('5 Proactivity & Clarity — 主动引导与清晰性','预先提供线索，让人知道可以做什么及当前状态。','靠近发光、轮次提示、占用昵称提示。','按Session 5第10页区分affordance（能做什么）与signifier（怎样传达）。'),('6 Ethics & Inclusion — 伦理与包容','考虑可访问性、安全、隐私、同意及不同参与方式。','临时昵称身份、无需永久账号、静音/退出、可观察与聆听、8秒让位。','手机/网络依赖、语言、视听能力和公开文字内容仍需改进；不宣称已全面无障碍。')]
t=d.add_table(rows=1,cols=4);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['原则','课堂含义','项目证据','限制/理论说明']):c.text=v
for row in principles:
 for c,v in zip(t.add_row().cells,row):c.text=v
p('理论使用方式：引用课堂课件，再解释自己的设计。CSCW来自Session 2第8页；affordances、signifiers、discoverability和mapping来自Session 5第10页。没有阅读原著时，不要假装直接阅读并引用原著；写为课件中讨论的理论，或按unit outline规则使用secondary citation。')
h('七、A1变化与反馈证据')
p('当前讲稿只使用开发记录可以支持的变化：多人同步、移动视角、钓鱼难度、钢琴面板与8秒让位、删除路灯占位。你提供的Social Interactive.pptx第5页提出日常生活/文化提示与共享视觉输出，第7页提出共享地图/角色/合作活动，第8页提出代币想法。最终版本保留跨文化连接目标、二维码/手机/大屏及角色地图，转向自由探索和短活动，没有加入代币。若此文件只是概念草稿，请与正式A1提交版核对。若有导师或同学反馈，请提供真实记录；没有正式用户研究就明确说开发测试和设备反馈。')
h('八、AI声明草稿（核实后再提交）')
p('I used ChatGPT/Codex for assistance with parts of the prototype code, debugging, iteration planning and presentation drafting. Image-generation tools were used to generate or refine selected activity icons and fish artwork/cut-outs. I selected and revised the visual direction, integrated outputs into the existing p5.js project, and tested/refined the interactions. The supplied map and character assets were retained; their original sources and licences are recorded separately. The final submission includes a description of where AI outputs were used and how they were adapted. [Confirm tutor approval and the unit’s permitted use before submission.]')
p('这段不能代替导师批准，也不能笼统声明“全部原创”或“仅用于润色”。按真实使用情况补工具名称、日期、输出用途和你自己的修改。保留提示词、版本截图或开发记录，以便说明作者决策。')
h('九、引用与素材来源')
p('PPT目前采用临时作者-年份格式。Brief没有指定具体引用格式，只说使用unit outline规定的格式，因此请提供unit outline后统一调整。原图作者未知的地方要补来源，不要写“我手绘”或“无版权”。')
for v in ['GRDE2053 Social Interactive Experiences. (n.d.). Assignment brief: Exhibition showcase [Course handout; supplied as A2_Brief.pdf].','GRDE2053 Social Interactive Experiences. (n.d.). A2: Exhibition showcase marking rubric [Assessment rubric; supplied as SIE-DCI-A2-rubric.pdf].','GRDE2053 Social Interactive Experiences. (n.d.). Session 2 [Lecture slides], pp.6,8,10; Session 3 [Lecture slides], pp.8,12; Session 5: Designing Interactions [Lecture slides], slides6–10. Supplied course materials.', 'Social Interactive. (n.d.). A Social Interactive Experience for Campus Connection [User-supplied concept deck], slides5,7–8.','p5.js. https://p5js.org/','Firebase. Realtime Database documentation. https://firebase.google.com/docs/database/','GitHub. GitHub Pages. https://pages.github.com/','MDN Web Docs. Web Audio API. https://developer.mozilla.org/en-US/docs/Web/API/Web_Audio_API','Project screenshots and source: https://github.com/itsangelia/public-map-project (2026).']:p(v)
h('十、提交前最后核对')
for v in ['更换封面的[Your name]。','替换PPT7的视频占位，插入真实75秒视频，并测试播放音量。','核对PPT9的已确认原则及项目解释；确认初始概念与正式A1的一致性。','带手机完整排练一次，实测5–7分钟。','确认导师已批准AI用法，附真实声明和素材来源。','独立完成最多15页的development documentation PDF。','按brief准备5–7分钟完整展示录制MP4，包含1–2分钟showcase；上传SharePoint，在coversheet贴导师可访问的链接。若老师另有录制范围要求，以老师确认为准。','将coversheet、过程PDF及要求的录制材料整理进ZIP；检查SharePoint权限与Week10记录。']:p('□ '+v)
d.save(O/'Teacher-requirements-and-rehearsal-ZH.docx')
rows=[['asset_or_media','use','known_source','action_before_submission'],['Prototype screenshots','PPT1,5,7','Controlled browser screenshots from project repository,2026','Do not describe as participant-study evidence'],['map draft.png','Map background','User-supplied; original creator not documented','Confirm creator,original source and usage rights'],['男角色.png / 女角色.png / 购物袋.png','Avatars','User-supplied; original creator not documented','Confirm creator,original source and usage rights'],['drink-shop.png / fishing-spot.png / shared-piano.png / wishing-tree.png','Activity icons','AI-generated/refined during project development','Declare tool,use and subsequent selection/refinement'],['Fish originals / cut-outs','Catch results','Some user-supplied; AI-assisted extraction/background removal; other fish generated','Record original sources and AI steps'],['Live piano tones','Audio','Synthesized by project Web Audio implementation','Declare code assistance; no external music recording claimed'],['Course PDFs','Assignment and rubric','User-supplied course materials','Add course-author/year if supplied'],['Six Social UX principles','Theory','Session2 p6;Session3 p8;Session5 slides6–7, supplied course materials','Exact labels verified;retain lecture citations'],['Showcase video','PPT7','Not supplied','Record real participant interaction;insert before presentation']]
with (O/'Media-and-AI-source-register.csv').open('w',newline='',encoding='utf-8-sig') as f:csv.writer(f).writerows(rows)
readme='''# A2 Presentation Package — Course-verified revision\n\nPrepared from the two supplied teacher PDFs and the actual project development record.\n\n## Full feature coverage
QR / nickname / three avatars / shared map; five-second chat; wish notes, previews and blessings; alternating drink creation; split-axis fishing and four catches; on-map shared piano and automatic release.

## Main files\n- Shared-Campus-A2.pptx — editable15-slide PowerPoint with bilingual presenter notes.\n- Shared-Campus-A2.pdf — matching visual review copy.\n- Speech-EN-with-Chinese-notes.docx — timed English script with Chinese explanation.\n- Teacher-requirements-and-rehearsal-ZH.docx — rubric mapping,video shot list,demo steps,AI draft and source checklist.\n- Media-and-AI-source-register.csv — attribution items to confirm.\n\n## Not yet presentation-ready\nThe six course principle labels are now verified against Sessions2,3 and5. The initial concept comparison uses Social Interactive.pptx. Slide7 still needs a real showcase video,and the cover needs the presenter name. Confirm whether the concept deck is the actual A1 submission;asset provenance and tutor AI approval also need confirmation. No participant numbers,formal research results or tutor feedback were invented.\n\nTarget total runtime6:55 including75 seconds of video and60 seconds live demonstration. Appendix slide15 is not part of the timed talk. Verify with a real rehearsal.\n\nThis package is not the complete assignment submission: the separate development document(max15 pages excluding title/references),coversheet,recorded MP4/SharePoint access and AI declaration are still required by the brief.\n\nNo project files or deployed game behavior were modified.\n'''
(O/'READ-ME-FIRST.md').write_text(readme)
print('Support documents created.')
