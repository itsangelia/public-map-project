from pathlib import Path
import re,json,math,zipfile
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Pt as DPt
O=Path('/workspace/presentation-complete');P=Path('/workspace/public-map-project');A=P/'previews'
BG='#FFF8F1';NAVY='#303A68';INK='#4D4B46';SAGE='#E0E6CE';PINK='#EFDAD3';WOOD='#DABDA5';MUTED='#76746E';WHITE='#FFFFFF'
slides=[
 dict(title='Shared Campus',time='0:00–0:20',seconds=20,en="Hello, today I’m presenting Shared Campus: a multiplayer campus map designed to make a first connection feel less awkward. People join on their own phones, see one another on a public screen, and choose a small activity to share.",zh='大家好，今天展示的是 Shared Campus：一个希望降低初次交流压力的多人校园地图。参与者用自己的手机加入，在公共大屏上看见彼此，并选择一个简短的共同活动。',cue='介绍项目，不要逐个列举功能。填上你的名字。'),
 dict(title='A small invitation to connect',time='0:20–0:55',seconds=35,en="My social focus is connection, especially the first interaction between international students who may not already know each other. Instead of asking strangers to begin a conversation, I offer something they can do together. The map also supports joy through playful outcomes, and reflection through shared wishes. These are design intentions, rather than measured social outcomes.",zh='核心社会目标是连接，特别是不熟悉彼此的国际学生之间的第一次互动。我没有要求陌生人先开口聊天，而是提供可以共同完成的小事。趣味结果带来快乐，许愿树提供表达与反思的空间。这些是设计目标，不代表已经测量并证明能改善孤独感或归属感。',cue='清楚说出 Connection 是主要焦点；不要把设计动机说成调研结论。'),
 dict(title='One shared place. Your own way in.',time='0:55–1:20',seconds=25,en="The entry flow is simple: scan the code, choose a nickname and character, and explore. The public display makes other participants visible, while each phone follows its own character. A first encounter is designed to fit within about two minutes. People can collaborate, leave a wish, listen, or observe before joining.",zh='入口流程很简单：扫码、选择昵称和角色、进入地图。公共大屏让参与者看见彼此，手机则跟随自己的角色。第一次简短互动以约两分钟内为目标。参与者可以合作、留言许愿、听琴，也可以先观察。',cue='指出“两分钟”指一次简短参与，不是要求在两分钟内体验完全部地图。'),
 dict(title='See the experience in use',time='1:20–2:35',seconds=75,en="This short film shows how people enter the shared space and move from noticing each other to participating together.",zh='这段短片展示参与者如何进入共享空间，并从注意到彼此，走向共同参与。',cue='播放你实拍的75秒展示视频。当前文件没有嵌入真实参与者视频，必须替换此页占位。不要朗读拍摄分镜。'),
 dict(title='A live moment of cooperation',time='2:35–3:35',seconds=60,en="Now I’ll demonstrate the drink activity with two phones. Each person makes alternating choices, and both receive the drink they created together.",zh='现在用两台手机现场演示饮品制作。两人轮流选择，最后共同得到制作完成的饮品。',cue='提前让两台手机加入 showcase 房间并到饮品店附近。进入匹配→A选基底→B选风味→A选配料→B选装饰→同时展示同一饮品结果→返回地图。控制在60秒内。'),
 dict(title='The six Social UX principles in practice',time='3:35–4:25',seconds=50,en="Here I use the course’s six principles. Sensory Feedback appears in piano sound. Embodied Interaction connects dragging to movement. Social Dynamics comes from shared presence and split roles. Resilience and Motivation shaped easier fishing and retry. Proactivity and Clarity appear in proximity glow and turn messages. Ethics and Inclusion informed temporary identities, mute and clear exits. The lectures’ CSCW perspective explains the shared roles; affordances and signifiers explain how people discover actions. However, phone and network requirements still exclude some visitors.",zh='这里使用课堂中六项原则。感官反馈体现在钢琴声音；具身互动体现在拖动与移动的对应；社会动态来自共同在场与角色分工；韧性与动机促成更宽松的钓鱼和重试；主动引导与清晰性体现在靠近发光及轮次提示；伦理与包容影响了临时身份、静音和退出设计。课件的CSCW视角帮助解释共享角色，affordances和signifiers解释人们怎样发现可做的事。不过，手机和网络要求仍会排除部分访客。',cue='六项名称已核实：Session 2 PDF第6页、Session 3 PDF第8页；Session 5第6–7页也列出同一框架。CSCW来自Session 2第8页；affordances/signifiers来自Session 5第10页。具身互动仅是手机触控与映射，不要夸大为全身沉浸；未实现触觉反馈。伦理设计不等于全面可访问，公共昵称与愿望仍需隐私、内容和可访问性反思。'),
 dict(title='A campus that feels inviting',time='4:25–4:50',seconds=25,en="Rather than a linear story, the experience uses a familiar campus setting as its narrative frame. Warm colours, illustrated characters and consistent activity icons create a welcoming atmosphere. Nearby objects glow to invite interaction. Shared drinks, unexpected fish and music provide small sensory rewards, while the public screen makes these individual actions part of a collective scene.",zh='体验没有使用线性剧情，而是把熟悉的校园作为叙事框架。暖色、插画角色和统一活动图标营造友好的氛围。靠近时发光的入口引导探索；饮品、意外的鱼和音乐带来小小的感官回报，公共大屏则把个人行动呈现在同一个场景里。',cue='结合截图说明视觉和声音怎样服务体验，不只说“好看”。'),
 dict(title='Refinement came from specific problems',time='4:50–5:30',seconds=40,en="My initial concept proposed cultural prompts, shared visuals and tokens. The final version keeps the connection goal but uses an explorable map and short activities, without tokens. Device feedback led to clearer presence, a larger phone view and gentler fishing. The piano became an on-map panel, with eight-second inactivity release. I also removed the unused streetlight placeholder.",zh='最初概念提出文化相关提示、共享视觉结果和代币。最终版本保留连接目标，但改为可探索地图与短活动，没有加入代币。设备反馈促成更清楚的共同在场、更大的手机视角与更宽松的钓鱼。钢琴变为地图上的操作面板，8秒无音符后自动让位；未使用的路灯占位被移除。',cue='初始概念依据Social Interactive.pptx第5、7、8页：日常生活/文化提示、共享地图、代币想法。若该文件不是正式A1提交版，展示前与提交版核对；不宣称这些是导师要求改动。技术迭代依据开发记录和设备反馈，不虚构正式访谈或观众人数。'),
 dict(title='Technology supports the social experience',time='5:30–6:00',seconds=30,en="The original p5.js project remains the foundation. Firebase shares positions and activity state; GitHub Pages hosts the site. Phones provide personal control, and the screen provides a common view. I checked simulated clients, browsers and actual devices, including iPhone sound. Exhibition still needs stable internet, charged devices and a recorded backup.",zh='原有p5.js项目仍然是基础。Firebase同步玩家位置和活动状态，GitHub Pages托管页面。手机负责个人操作，大屏提供共同视角。开发检查包括独立模拟客户端与浏览器测试；实际手机测试也在修复后确认了iPhone出声。展览时仍需准备稳定网络、充足电量和录像备份。',cue='不要把26项自动测试说成26位用户；不要声称在所有设备、所有网络环境下都稳定。'),
 dict(title='A complete prototype — with honest limits',time='6:00–6:35',seconds=35,en="People can watch, listen, leave a wish or cooperate. The value is in these small invitations. I used generative AI for parts of the code, debugging, artwork and presentation, then selected, tested and refined its outputs. Next, I would observe whether unfamiliar participants can begin an interaction independently, rather than adding more activities.",zh='最终体验提供不同程度的社交投入：观看、聆听、留下愿望，或共同合作。价值不在功能数量，而在这些功能创造的小邀请。我在部分代码、调试、活动素材及演示准备中使用了生成式AI，再进行选择、测试与调整。下一步更值得观察不熟悉的参与者是否能独立开始有意义的互动，而不是继续堆加功能。',cue='务必核实并补齐AI来源、导师批准和单元允许范围。不要把AI辅助内容宣称为完全独立编写或手绘。'),
 dict(title='Thank you',time='6:35–6:55',seconds=5,en="Thank you. I’d like to invite you to try the map together.",zh='谢谢，欢迎大家一起体验这个地图。',cue='参考资料页保留在屏幕上；不逐条朗读参考文献。'),
 dict(title='Submission checklist',time='APPENDIX — do not present',seconds=0,en='',zh='',cue='备查页，不计入课堂展示。按brief还需封面、独立development documentation PDF、MP4及SharePoint可访问链接；这套PPT不替代过程文档。')]
# Full feature coverage before the filmed experience.
extra=[
 dict(title='Meet, speak and leave a little hope',seconds=30,en="Players choose from three characters and see each other moving independently. Chat appears above their heads and clears after five seconds. At the wishing tree, anyone can hang a wish. Others see a short, swaying preview, open the full note with its nickname and date, and add a blessing.",zh='玩家选择三种角色，在同一地图里独立移动、看见彼此。聊天气泡显示在头顶，5秒后消失。许愿树支持挂愿望；其他人可看简短摆动预览，点开完整内容、昵称与日期，并送上祝福。',cue='强调聊天与许愿是两种不同的表达：即时交流与保留在树上的留言。纸条仅预览4–5个英文词；每人对同一愿望祝福一次。'),
 dict(title='Two ways to make something together',seconds=30,en="The drink shop matches two named players. They alternate four choices, with three options at each stage, then receive the same named drink. Fishing requires a different cooperation: one player moves the float vertically, the other horizontally. Staying over the ripples builds progress. Success reveals one of four named fish; failure offers another try.",zh='饮品店匹配两名显示昵称的玩家，四道程序轮流选择，每道三个选项，共同得到同一杯有名字的饮品。钓鱼则由一人控制上下、一人控制左右，对准涟漪积累进度。成功随机获得四种有名字的鱼之一，失败可以重试。',cue='饮品：基底/风味/配料/装饰，81种组合。钓鱼：倒计时后开始，20%初始进度，对准+10%/秒，偏离−4%/秒，降至0失败。目标3.5秒换位，容错范围扩大10%。'),
 dict(title='One performer. A shared audience.',seconds=25,en="The piano stays within the map. One person plays twenty-one keys while others listen. A moving icon signals the performance, and tapping an occupied piano shows the performer’s name. After eight seconds without a note, the seat becomes available. Muting, clear exits and proximity cues let visitors choose how much to participate.",zh='钢琴以地图内面板打开。一人演奏21个音键，其他人聆听。图标动态提示正在演奏；点击被占用的钢琴会显示演奏者昵称。8秒不按音符自动让位。静音、明确退出与靠近提示支持不同参与程度。',cue='地图仍可见；旁观者不会打开同步键盘。靠近活动入口会出现提醒。说明观看、听琴也属于参与方式。')]
slides[3:3]=extra
# Budget: 6:55 including film and live interaction.
budgets=[15,20,15,30,30,25,75,60,45,15,25,20,30,10,0]
clock=0
for n,(m,secs) in enumerate(zip(slides,budgets)):
 m['seconds']=secs
 if secs:
  m['time']=f'{clock//60}:{clock%60:02d}–{(clock+secs)//60}:{(clock+secs)%60:02d}'
  clock+=secs
 else:m['time']='APPENDIX — do not present'
# Avoid repeating the feature inventory in later reflection.
slides[0]['en']="Hello, I’m presenting Shared Campus, a multiplayer campus map that offers small, low-pressure ways for people to connect through their phones and a public display."
slides[1]['en']="My primary focus is connection, particularly between international students who do not yet know each other. A shared action can make the first interaction easier. Playful outcomes support joy, while wishes invite reflection. These are design intentions, not measured social outcomes."
slides[2]['en']="Visitors scan a QR code, choose a nickname and avatar, then explore. Each phone follows its own player; the public screen shows the shared world. One short encounter is designed to fit within about two minutes."
slides[9]['en']="A familiar illustrated campus provides the narrative setting. Soft colours, consistent icons, proximity glow and sound connect the different activities into one welcoming place."
slides[10]['en']="My initial concept used cultural prompts, shared visuals and tokens. I kept the connection goal, but moved to free exploration and short activities without tokens. Device feedback shaped clearer multiplayer presence, thirty-five percent phone zoom, gentler fishing and the piano’s automatic release."
slides[11]['en']="p5.js renders the world, Firebase shares state, and GitHub Pages hosts it. Checks included simulated clients, browsers and actual phone feedback. Reliable internet and a recorded fallback remain necessary for exhibition."
# Native editable PowerPoint + matching rendered PDF/previews.
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
W,H=1600,900; canvases=[];manifest=[]
F='/usr/share/fonts/truetype/dejavu/'
def font(sz,bold=False,serif=False):return ImageFont.truetype(F+('DejaVuSerif' if serif else 'DejaVuSans')+('-Bold' if bold else '')+'.ttf',sz)
def rgb(c):return RGBColor.from_string(c.strip('#'))
def rect(x,y,w,h,color,r=0,line=None):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x/120),Inches(y/120),Inches(w/120),Inches(h/120));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=rgb(line)
 if r:
  try:sh.adjustments[0]=.14
  except:pass
 draw.rounded_rectangle((x,y,x+w,y+h),r,fill=color,outline=line) if r else draw.rectangle((x,y,x+w,y+h),fill=color,outline=line)
 return sh
def wrap(txt,w,sz,bold=False,serif=False):
 f=font(sz,bold,serif);out=[]
 for para in txt.split('\n'):
  words=para.split();line=''
  for word in words:
   test=(line+' '+word).strip()
   if f.getlength(test)>w and line:out.append(line);line=word
   else:line=test
  out.append(line)
 return '\n'.join(out)
def text(x,y,w,txt,sz=34,color=INK,bold=False,serif=False):
 txt=wrap(txt,w,sz,bold,serif);lines=txt.split('\n');lineh=sz*1.28
 box=sl.shapes.add_textbox(Inches(x/120),Inches(y/120),Inches(w/120),Inches((len(lines)*lineh+8)/120));tf=box.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Georgia' if serif else 'Arial';p.font.size=Pt(sz*.6);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(0);p.line_spacing=1.05
  draw.text((x,y+i*lineh),line,font=font(sz,bold,serif),fill=color)
 return box

def pic(path,x,y,w,h):
 path=Path(path);im=Image.open(path).convert('RGBA');ratio=min(w/im.width,h/im.height);iw,ih=im.width*ratio,im.height*ratio;xx,yy=x+(w-iw)/2,y+(h-ih)/2
 sl.shapes.add_picture(str(path),Inches(xx/120),Inches(yy/120),width=Inches(iw/120),height=Inches(ih/120));canvas.alpha_composite(im.resize((round(iw),round(ih)),Image.Resampling.LANCZOS),(round(xx),round(yy)))
def card(x,y,w,h,num,title,body,color=SAGE):
 rect(x,y,w,h,color,18);text(x+22,y+17,w-44,num,20,NAVY,True);text(x+22,y+50,w-44,title,30,NAVY,True);text(x+22,y+50+len(wrap(title,w-44,30,True).split('\n'))*38.4+15,w-44,body,25)
def head(kicker,title):text(75,48,1450,kicker.upper(),20,MUTED,True);text(75,93,1450,title,54,NAVY,True,True)
for page_index,meta in enumerate(slides):
 i=page_index if page_index<3 else page_index-3
 if 3<=page_index<=5:i=100+page_index
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(BG);canvas=Image.new('RGBA',(W,H),BG);draw=ImageDraw.Draw(canvas)
 text(75,850,1300,'SHARED CAMPUS  /  GRDE2053 Social Interactive Experiences',17,MUTED);text(1470,850,70,f'{page_index+1:02d}',20,NAVY,True)
 if i==0:
  pic(A/'fishing-map-overview.png',620,75,920,740);rect(0,0,620,830,BG);text(75,92,510,'EXHIBITION SHOWCASE',20,MUTED,True);text(75,180,530,'Shared\nCampus',87,NAVY,True,True);text(75,445,480,'Small moments.\nShared connections.',38,NAVY);text(75,610,480,'A public map, personal phones\nand playful first encounters.',27);text(75,735,510,'[Your name]  ·  A2  ·  2026',23,MUTED)
 elif i==1:
  head('01 / Social focus','A small invitation to connect');text(80,237,840,'How might a shared campus space make a first interaction easier?',44,NAVY,True)
  card(80,415,450,260,'PRIMARY','Connection','A small shared action before a conversation.');card(575,415,450,260,'SECONDARY','Joy','Playful choices and unexpected outcomes.',PINK);card(1070,415,450,260,'ALSO','Reflection','Leave a wish. Read someone else’s.',WOOD);text(80,738,1400,'Audience focus: international students and other campus visitors.',28,MUTED)
 elif i==2:
  head('02 / Participant journey','One shared place. Your own way in.');steps=[('01','Scan','Use your own phone.'),('02','Choose','Nickname + avatar.'),('03','Explore','See others nearby.'),('04','Notice','Follow a glowing invitation.'),('05','Share','Make, play, wish or listen.')]
  for j,(n,t,b) in enumerate(steps):card(75+j*294,255,278,260,n,t,b,SAGE if j%2==0 else PINK)
  rect(75,580,1450,155,NAVY,20);text(110,603,1370,'A first encounter in approximately two minutes or less.',38,WHITE,True);text(110,663,1370,'The whole map can be explored at your own pace.',27,'#EBE6DD')
 elif i==103:
  head('Experience / expression','Meet, speak and leave a little hope')
  pic(P/'wishing-tree.png',980,220,510,525)
  card(75,235,430,220,'JOIN + MOVE','Be present','QR entry, nickname, three avatars; independent movement.')
  card(540,235,430,220,'LIVE CHAT','Speak in the map','Shared head bubbles disappear after 5 seconds.',PINK)
  card(75,490,895,255,'WISHING TREE','A message someone can discover','Hang a wish. A 4–5-word preview sways; open the note for nickname, date and full text. Bless it once per player.',WOOD)
 elif i==104:
  head('Experience / cooperation','Two players. Two kinds of cooperation.')
  pic(A/'drink-maker-online-finished.png',60,220,245,530)
  text(330,240,420,'Make a drink',36,NAVY,True)
  text(330,305,420,'Named player matching\nFour alternating turns\nThree options per stage\nYour turn notification\nA shared, named result',26)
  text(330,565,420,'Base → flavour → topping → decoration',25,NAVY,True)
  pic(A/'fishing-online-a.png',800,220,245,530)
  text(1075,240,440,'Catch together',36,NAVY,True)
  text(1075,305,440,'Countdown to start\nSplit vertical / horizontal control\nAim the float at ripples\nProgress rises or falls\nFour named catches + retry',26)
  text(1075,590,440,'Start 20% · +10% / −4% per second',24,NAVY,True)
  text(80,780,1440,'Drink choices: sparkling water / milk tea / slush; strawberry / matcha / mango; pearls / coconut jelly / red bean; foam / ice cream / lemon.',18,MUTED)
 elif i==105:
  head('Experience / music + shared access','One performer. A shared audience.')
  pic(A/'piano-playing.png',85,215,290,575)
  pic(P/'shared-piano.png',1150,250,340,360)
  text(440,250,680,'Play without leaving the map',38,NAVY,True)
  text(440,335,650,'21 keys in an on-map panel\nOne performer; others can listen\nOccupied notice shows the nickname\nPlaying icon moves with musical notes\n8 seconds idle → seat released\nSound control and a clear exit',28)
  rect(440,665,1030,105,SAGE,18)
  text(465,687,975,'All activity entrances share proximity cues and a route back to the map.',27,NAVY)
 elif i==3:
  head('03 / Showcase film','From noticing to participating');rect(75,215,1450,535,NAVY,24);text(145,278,1300,'INSERT YOUR SHOWCASE VIDEO',38,WHITE,True);text(145,355,1270,'75 seconds  /  two participants + phones + public display',30,'#EAE4D7');text(145,455,1180,'Join → meet → collaborate → reflect → listen',42,WHITE,True);text(145,585,1200,'Replace this page with real recorded footage before presenting.',26,'#EFDAD3');text(80,790,1400,'Brief requirement: a 1–2 minute showcase video of the experience in use.',25,MUTED)
 elif i==4:
  head('04 / Live demonstration','Two people. Four choices. One shared drink.');
  for j,n in enumerate(['drink-maker-matching.png','drink-maker-your-turn.png','drink-maker-online-finished.png']):pic(A/n,75+j*355,230,310,550)
  pic(O/'join-qr.png',1190,270,290,290);text(1170,600,330,'LIVE DEMO',23,NAVY,True);text(1170,640,330,'Room: showcase\nTwo phones, one display\nTarget: 60 seconds',25);b=text(1170,760,330,'Open public display ↗',23,NAVY,True);b.click_action.hyperlink.address='https://itsangelia.github.io/public-map-project/campus-f3cd85c8e3d0eaa9.html?mode=display&room=showcase'
 elif i==5:
  head('05 / Social UX reflection','The six course principles, applied');text(80,185,1400,'A course framework connected to concrete design decisions.',24,MUTED)
  examples=[('01','Sensory Feedback','Piano sound confirms actions.'),('02','Embodied Interaction','Dragging maps to movement.'),('03','Social Dynamics','Co-presence and split roles.'),('04','Resilience & Motivation','Gentler fishing and retry.'),('05','Proactivity & Clarity','Proximity glow and turn cues.'),('06','Ethics & Inclusion','Temporary IDs, mute and exits.')]
  for j,(n,t,b) in enumerate(examples):card(75+j%3*493,250+j//3*255,470,225,n,t,b,SAGE if j%2==0 else PINK)
  text(80,778,1400,'Framework: Session 2, p. 6; Session 3, p. 8. Grounding: CSCW / affordances / signifiers (Sessions 2 & 5).',19,MUTED)
 elif i==6:
  head('06 / Atmosphere + sensory experience','A familiar campus, with small invitations');pic(A/'piano-playing.png',1020,205,300,610)
  for j,(t,b) in enumerate([('A narrative setting','A shared campus rather than a linear quest.'),('Visual invitation','Consistent icons, soft colours and proximity glow.'),('Sensory reward','A shared result, a surprising fish, a live melody.')]):text(80,245+j*175,850,t,35,NAVY,True);text(80,302+j*175,825,b,30)
 elif i==7:
  head('07 / Development + refinement','From initial concept to refined interactions');rows=[('Concept shift','Cultural prompts / shared visuals','Free exploration + shared activities'),('Reliability','Missing players / messages','Refined presence + shared state'),('Mobile view','Small avatar / unused space','35% zoom + full viewport'),('Cooperation','Jerky / difficult fishing','Smoothing + gentler timing'),('Shared access','Separate piano page / idle use','Map controls + 8-second release')]
  for j,(area,issue,change) in enumerate(rows):
   y=230+j*105;rect(75,y,1450,88,SAGE if j%2==0 else '#F3E9DE',13);text(98,y+22,235,area,27,NAVY,True);text(350,y+22,510,issue,26);text(900,y+22,590,change,26,NAVY)
  text(80,785,1400,'Initial concept: Social Interactive, slides 5, 7–8. Tokens not carried forward; unused streetlight placeholder removed.',21,MUTED)
 elif i==8:
  head('08 / Technical implementation','A small architecture for shared participation');card(75,245,430,245,'PERSONAL INPUT','Phones','Choose, move and interact.');card(585,245,430,245,'SHARED STATE','Firebase','Positions, turns, notes and wishes.',PINK);card(1095,245,430,245,'COMMON VIEW','Public display','See co-presence and shared outcomes.',WOOD);text(510,345,70,'↔',42,NAVY);text(1020,345,70,'↔',42,NAVY)
  text(80,560,1380,'p5.js renders the world. GitHub Pages hosts the site.',35,NAVY,True);text(80,635,1380,'Verification: 26 local checks + browser checks + reported device trials.',29);text(80,700,1380,'Exhibition needs: stable internet, charged phones and a recorded fallback.',29);text(80,765,1380,'Automated checks are not a substitute for audience evaluation.',23,MUTED)
 elif i==9:
  head('09 / Reflection + authorship','Small invitations. Shared experiences.');text(80,235,900,'Watch. Listen. Leave a wish.\nOr do something together.',47,NAVY,True,True);text(80,415,900,'Next evaluation:\nCan unfamiliar participants begin independently?',34);text(80,585,885,'AI-assisted development was selected, tested and refined — not treated as finished output.',31)
  rect(1040,230,475,515,SAGE,22);text(1070,260,410,'AI USE',23,NAVY,True);text(1070,320,400,'Code + debugging\nSome activity artwork\nPresentation drafting',31);text(1070,515,400,'Confirm tutor approval, permitted use and the required declaration before submission.',25);text(80,765,1400,'Impact on belonging has not yet been established through an independent participant study.',22,MUTED)
 elif i==10:
  head('References + credits','Thank you. Try it together.');text(80,235,1380,'Course sources',31,NAVY,True);text(80,290,1380,'GRDE2053. A2 assignment brief, pp. 1–5. Supplied course document, n.d.\nGRDE2053. SIE–DCI A2 Exhibition Showcase marking rubric. Supplied course document, n.d.\nGRDE2053. Sessions 2 (p. 6), 3 (p. 8) and 5 (slides 6–10). Course slides, n.d.\nInitial concept: Social Interactive.pptx, slides 5, 7–8. User-supplied concept deck, n.d.',25)
  text(80,490,1380,'Project media + implementation',31,NAVY,True);text(80,545,1380,'Prototype screens: public-map-project repository, 2026; controlled browser captures.\nMap / avatars: user-supplied assets; confirm original creators and licences.\nSome activity icons / fish cut-outs: AI-generated or AI-refined; list individual assets.\np5.js • Firebase Realtime Database • GitHub Pages • Web Audio API.',25)
  text(80,725,1380,'References here use a provisional author–date format. Match the referencing style in your unit outline.',22,MUTED)
 else:
  head('Appendix / not part of the timed talk','Before you submit');items=['Fill in your name; check the verified theory mapping and initial-concept comparison.','Record and insert a real 1–2 minute showcase film.','Rehearse a 5–7 minute presentation including video and live demo.','Prepare the separate development PDF: maximum 15 pages, excluding title / references.','Include coversheet, requested MP4 / SharePoint link and tutor access.','Confirm AI approval, declaration and media attribution; check milestone completion.']
  for j,t in enumerate(items):text(85,235+j*85,1410,f'{j+1:02d}  {t}',29,NAVY if j<3 else INK)
 notes=f"SLIDE {page_index+1} | {meta['time']}\n\nENGLISH SCRIPT\n{meta['en']}\n\n中文理解版\n{meta['zh']}\n\nPRESENTER CUE\n{meta['cue']}"
 sl.notes_slide.notes_text_frame.text=notes
 if i not in [0,10,11]:text(80,820,1400,'Images: prototype assets / controlled browser screenshots, project repository (2026).',14,MUTED)
 path=O/f'slide-{page_index+1:02d}.png';canvas.convert('RGB').save(path);canvases.append(canvas.convert('RGB'));manifest.append({'slide':page_index+1,**meta})
prs.save(O/'Shared-Campus-A2.pptx')
canvases[0].save(O/'Shared-Campus-A2.pdf',save_all=True,append_images=canvases[1:],resolution=120)
(O/'presentation-content.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
# Script, bilingual explanation and timings.
doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=DPt(11)
doc.add_heading('Shared Campus — Presentation Script',0);doc.add_paragraph('GRDE2053 • A2 Exhibition Showcase • Draft for review • English delivery with Chinese explanation')
doc.add_paragraph('Target runtime: 6:55 including a75-second showcase video and60-second live demonstration. Do not read the Chinese paragraphs aloud. Teacher documents require both video and live demonstration. The theory mapping is verified against the supplied course slides. Insert your real showcase video and name before presenting.')
for i,m in enumerate(slides[:-1]):
 doc.add_heading(f"{i+1:02d}  {m['title']}  |  {m['time']}",1)
 doc.add_paragraph(m['en']);doc.add_paragraph('中文理解：'+m['zh']);doc.add_paragraph('操作提示：'+m['cue'])
doc.add_heading('Timing and rehearsal',1);doc.add_paragraph('Use a timer. Aim for clear conversational delivery rather than reading every word. If the full rehearsal exceeds7 minutes, shorten technical detail and reflection; keep the required video, live demonstration and course-principle reflection. Rehearse using your actual phones and network.')
doc.add_heading('Important evidence boundaries',1);doc.add_paragraph('Development feedback and user confirmations in the project log are not a formal research study. Do not invent participant numbers, interview quotes, tutor feedback or A1 claims. The six principles are verified in Session 2 p.6 and Session 3 p.8. The initial concept comparison uses Social Interactive.pptx; confirm whether it is the actual A1 submission. Slide7 still needs real showcase footage. AI-assisted outputs must be disclosed according to tutor-approved unit rules.')
doc.save(O/'Speech-EN-with-Chinese-notes.docx')
md='# Shared Campus — 逐页演讲稿\n\n目标6:55，包含75秒视频和60秒现场演示。中文仅供理解，不要两种语言各读一遍。\n\n'
for i,m in enumerate(slides[:-1]):md+=f"## {i+1}. {m['title']} | {m['time']}\n\n{m['en']}\n\n**中文理解：** {m['zh']}\n\n**操作提示：** {m['cue']}\n\n"
(O/'Speech-EN-ZH.md').write_text(md)
print('Created editable PPTX, matching PDF, bilingual DOCX and15 slide previews.')
print('Spoken English words:',sum(len(re.findall(r"\b[\w’–-]+\b",m['en'])) for m in slides))
