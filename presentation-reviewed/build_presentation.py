from pathlib import Path
import re,json,math,zipfile
from PIL import Image,ImageDraw,ImageFont
from pptx import Presentation
from pptx.util import Inches,Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Pt as DPt
O=Path('/workspace/presentation-reviewed');P=Path('/workspace/public-map-project');A=P/'previews'
BG='#FFF8F1';NAVY='#303A68';INK='#4D4B46';SAGE='#E0E6CE';PINK='#EFDAD3';WOOD='#DABDA5';MUTED='#76746E';WHITE='#FFFFFF'
slides=json.loads((O/'presentation-content.json').read_text())
# Native editable PowerPoint + matching rendered PDF/previews.
prs=Presentation();prs.slide_width=Inches(13.333333);prs.slide_height=Inches(7.5)
W,H=1600,900; canvases=[];manifest=[]
F='/usr/share/fonts/truetype/dejavu/'
def font(sz,bold=False,serif=False):return ImageFont.truetype(F+('DejaVuSerif' if serif else 'DejaVuSans')+('-Bold' if bold else '')+'.ttf',sz)
def rgb(c):return RGBColor.from_string(c.strip('#'))
def rect(x,y,w,h,color,r=0,line=None):
 sh=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x/120),Inches(y/120),Inches(w/120),Inches(h/120));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=rgb(line)
 if r:
  try:sh.adjustments[0]=.14
  except:pass
 draw.rounded_rectangle((x,y,x+w,y+h),r,fill=color,outline=line) if r else draw.rectangle((x,y,x+w,y+h),fill=color,outline=line)
 return sh
def wrap(txt,w,sz,bold=False,serif=False):
 f=font(sz,bold,serif);out=[]
 for para in txt.split('\n'):
  words=para.split();line=''
  for word in words:
   test=(line+' '+word).strip()
   if f.getlength(test)>w and line:out.append(line);line=word
   else:line=test
  out.append(line)
 return '\n'.join(out)
def text(x,y,w,txt,sz=34,color=INK,bold=False,serif=False):
 txt=wrap(txt,w,sz,bold,serif);lines=txt.split('\n');lineh=sz*1.28
 box=sl.shapes.add_textbox(Inches(x/120),Inches(y/120),Inches(w/120),Inches((len(lines)*lineh+8)/120));tf=box.text_frame;tf.clear();tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False
 for i,line in enumerate(lines):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Georgia' if serif else 'Arial';p.font.size=Pt(sz*.6);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(0);p.line_spacing=1.05
  draw.text((x,y+i*lineh),line,font=font(sz,bold,serif),fill=color)
 return box

def pic(path,x,y,w,h):
 path=Path(path);im=Image.open(path).convert('RGBA');ratio=min(w/im.width,h/im.height);iw,ih=im.width*ratio,im.height*ratio;xx,yy=x+(w-iw)/2,y+(h-ih)/2
 sl.shapes.add_picture(str(path),Inches(xx/120),Inches(yy/120),width=Inches(iw/120),height=Inches(ih/120));canvas.alpha_composite(im.resize((round(iw),round(ih)),Image.Resampling.LANCZOS),(round(xx),round(yy)))
def card(x,y,w,h,num,title,body,color=SAGE):
 rect(x,y,w,h,color,18);text(x+22,y+17,w-44,num,20,NAVY,True);text(x+22,y+50,w-44,title,30,NAVY,True);text(x+22,y+50+len(wrap(title,w-44,30,True).split('\n'))*38.4+15,w-44,body,25)
def head(kicker,title):text(75,48,1450,kicker.upper(),20,MUTED,True);text(75,93,1450,title,54,NAVY,True,True)
for page_index,meta in enumerate(slides):
 i=page_index if page_index<3 else page_index-3
 if 3<=page_index<=5:i=100+page_index
 sl=prs.slides.add_slide(prs.slide_layouts[6]);sl.background.fill.solid();sl.background.fill.fore_color.rgb=rgb(BG);canvas=Image.new('RGBA',(W,H),BG);draw=ImageDraw.Draw(canvas)
 text(75,850,1300,'SHARED CAMPUS  /  GRDE2053 Social Interactive Experiences',17,MUTED);text(1470,850,70,f'{page_index+1:02d}',20,NAVY,True)
 if i==0:
  pic(A/'fishing-map-overview.png',620,75,920,740);rect(0,0,620,830,BG);text(75,92,510,'EXHIBITION SHOWCASE',20,MUTED,True);text(75,180,530,'Shared\nCampus',87,NAVY,True,True);text(75,445,480,'Small moments.\nShared connections.',38,NAVY);text(75,610,480,'A public map, personal phones\nand playful first encounters.',27);text(75,735,510,'[Your name]  ·  A2  ·  2026',23,MUTED)
 elif i==1:
  head('Social focus','A small invitation to connect');text(80,237,840,'How might a shared campus space make a first interaction easier?',44,NAVY,True)
  card(80,415,450,260,'PRIMARY','Connection','A small shared action before a conversation.');card(575,415,450,260,'SECONDARY','Joy','Playful choices and unexpected outcomes.',PINK);card(1070,415,450,260,'ALSO','Reflection','Leave a wish. Read someone else’s.',WOOD);text(80,738,1400,'Audience focus: international students and other campus visitors.',28,MUTED)
 elif i==2:
  head('Participant journey','One shared place. Your own way in.');steps=[('01','Scan','Use your own phone.'),('02','Choose','Nickname + avatar.'),('03','Explore','See others nearby.'),('04','Notice','Follow a glowing invitation.'),('05','Share','Make, play, wish or listen.')]
  for j,(n,t,b) in enumerate(steps):card(75+j*294,255,278,260,n,t,b,SAGE if j%2==0 else PINK)
  rect(75,580,1450,155,NAVY,20);text(110,603,1370,'Exhibition target: join and share one short activity.',38,WHITE,True);text(110,663,1370,'Designed for a short encounter, with free exploration afterwards.',27,'#EBE6DD')
 elif i==103:
  head('Experience / expression','Meet, speak and leave a little hope')
  pic(P/'wishing-tree.png',980,220,510,525)
  card(75,235,430,220,'JOIN + MOVE','Be present','QR entry, nickname, three avatars; independent movement.')
  card(540,235,430,220,'LIVE CHAT','Speak in the map','Shared head bubbles disappear after 5 seconds.',PINK)
  card(75,490,895,255,'WISHING TREE','A message someone can discover','Hang a wish. A 4–5-word preview sways; open the note for nickname, date and full text. Bless it once per player.',WOOD)
 elif i==104:
  head('Experience / cooperation','Two players. Two kinds of cooperation.')
  pic(A/'drink-maker-online-finished.png',60,220,245,530)
  text(330,240,420,'Make a drink',36,NAVY,True)
  text(330,305,420,'Four alternating turns\nThree options per stage\nBoth players shape the recipe\nYour turn notification\nOne shared, named result',26)
  text(330,565,420,'Base → flavour → topping → decoration',25,NAVY,True)
  pic(A/'fishing-online-a.png',800,220,245,530)
  text(1075,240,440,'Catch together',36,NAVY,True)
  text(1075,305,440,'Countdown to start\nSplit vertical / horizontal control\nAim the float at ripples\nProgress rises or falls\nFour named catches + retry',26)
  text(1075,590,440,'Start 20% · +10% / −4% per second',24,NAVY,True)
  text(80,780,1440,'Drink choices: sparkling water / milk tea / slush; strawberry / matcha / mango; pearls / coconut jelly / red bean; foam / ice cream / lemon.',18,MUTED)
 elif i==105:
  head('Experience / music + shared access','One performer. A shared audience.')
  pic(A/'piano-playing.png',85,215,290,575)
  pic(P/'shared-piano.png',1150,250,340,360)
  text(440,250,680,'Play without leaving the map',38,NAVY,True)
  text(440,335,650,'21 keys in an on-map panel\nOne performer; others can listen\nOccupied notice shows the nickname\nPlaying icon moves with musical notes\n8 seconds idle → seat released\nSound control and a clear exit',28)
  rect(440,665,1030,105,SAGE,18)
  text(465,687,975,'All activity entrances share proximity cues and a route back to the map.',27,NAVY)
 elif i==3:
  head('Showcase film','From noticing to participating');rect(75,215,1450,535,NAVY,24);text(145,278,1300,'INSERT YOUR SHOWCASE VIDEO',38,WHITE,True);text(145,355,1270,'75 seconds  /  two participants + phones + public display',30,'#EAE4D7');text(145,455,1180,'Join → meet → collaborate → reflect → listen',42,WHITE,True);text(145,585,1200,'Replace this page with real recorded footage before presenting.',26,'#EFDAD3');text(80,790,1400,'Brief requirement: a 1–2 minute showcase video of the experience in use.',25,MUTED)
 elif i==4:
  head('Live demonstration','Two people. Four choices. One shared drink.');
  for j,n in enumerate(['drink-maker-matching.png','drink-maker-your-turn.png','drink-maker-online-finished.png']):pic(A/n,75+j*355,230,310,550)
  pic(O/'join-qr.png',1190,270,290,290);text(1170,600,330,'LIVE DEMO',23,NAVY,True);text(1170,640,330,'Room: showcase\nTwo phones, one display\nTarget: 55 seconds',25);b=text(1170,760,330,'Open public display ↗',23,NAVY,True);b.click_action.hyperlink.address='https://itsangelia.github.io/public-map-project/campus-f3cd85c8e3d0eaa9.html?mode=display&room=showcase'
 elif i==5:
  head('Social UX reflection','The six course principles, applied');text(80,185,1400,'A course framework connected to concrete design decisions.',24,MUTED)
  examples=[('01','Sensory Feedback','Piano sound responds to touch.'),('02','Embodied Interaction','Touchscreen control; limited embodiment.'),('03','Social Dynamics','Split controls make cooperation necessary.'),('04','Resilience & Motivation','Gentler fishing and retry.'),('05','Proactivity & Clarity','Proximity glow and turn cues.'),('06','Ethics & Inclusion','Choice and fair access; barriers remain.')]
  for j,(n,t,b) in enumerate(examples):card(75+j%3*493,250+j//3*255,470,225,n,t,b,SAGE if j%2==0 else PINK)
  text(80,778,1400,'Framework: Session 2, p. 6; Session 3, p. 8. Grounding: CSCW / affordances / signifiers (Sessions 2 & 5).',19,MUTED)
 elif i==6:
  head('Atmosphere + sensory experience','A familiar campus, with small invitations');pic(A/'piano-playing.png',1020,205,300,610)
  for j,(t,b) in enumerate([('A narrative setting','A shared campus rather than a linear quest.'),('Visual invitation','Consistent icons, soft colours and proximity glow.'),('Sensory reward','A shared result, a surprising fish, a live melody.')]):text(80,245+j*175,850,t,35,NAVY,True);text(80,302+j*175,825,b,30)
 elif i==7:
  head('Development + refinement','From initial concept to refined interactions');rows=[('Concept shift','Cultural prompts / shared visuals','Free exploration + shared activities'),('Reliability','Missing players / messages','Refined presence + shared state'),('Mobile view','Small avatar / unused space','35% zoom + full viewport'),('Cooperation','Jerky / difficult fishing','Smoothing + gentler timing'),('Shared access','Separate piano page / idle use','Map controls + 8-second release')]
  for j,(area,issue,change) in enumerate(rows):
   y=230+j*105;rect(75,y,1450,88,SAGE if j%2==0 else '#F3E9DE',13);text(98,y+22,235,area,27,NAVY,True);text(350,y+22,510,issue,26);text(900,y+22,590,change,26,NAVY)
  text(80,785,1400,'Initial concept: Social Interactive, slides 5, 7–8. Tokens not carried forward; unused streetlight placeholder removed.',21,MUTED)
 elif i==8:
  head('Technical implementation','A small architecture for shared participation');card(75,245,430,245,'PERSONAL INPUT','Phones','Choose, move and interact.');card(585,245,430,245,'SHARED STATE','Firebase','Positions, turns, notes and wishes.',PINK);card(1095,245,430,245,'COMMON VIEW','Public display','See co-presence and shared outcomes.',WOOD);text(510,345,70,'↔',42,NAVY);text(1020,345,70,'↔',42,NAVY)
  text(80,560,1380,'p5.js renders the world. GitHub Pages hosts the site.',35,NAVY,True);text(80,635,1380,'Checks: shared player visibility, messages, activity state and iPhone sound.',29);text(80,700,1380,'Exhibition needs: stable internet, charged phones and a recorded fallback.',29);text(80,765,1380,'Automated checks are not a substitute for audience evaluation.',23,MUTED)
 elif i==9:
  head('Reflection + authorship','Small invitations. Shared experiences.');text(80,235,900,'Watch. Listen. Leave a wish.\nOr do something together.',47,NAVY,True,True);text(80,415,900,'Next evaluation:\nDo shared activities prompt cultural exchange?',34);text(80,585,885,'AI-generated outputs were selected, tested and refined.',31)
  rect(1040,230,475,515,SAGE,22);text(1070,260,410,'AI USE',23,NAVY,True);text(1070,320,400,'Code + debugging\nSome activity artwork\nPresentation drafting',31);text(1070,515,400,'My decisions: visual selection, interaction refinement and device feedback.',25);text(80,765,1400,'Impact on belonging has not yet been established through an independent participant study.',22,MUTED)
 elif i==10:
  head('References + credits','Thank you. Try it together.');text(80,235,1380,'Course sources',31,NAVY,True);text(80,290,1380,'GRDE2053. A2 assignment brief, pp. 1–5. Supplied course document, n.d.\nGRDE2053. SIE–DCI A2 Exhibition Showcase marking rubric. Supplied course document, n.d.\nGRDE2053. Sessions 2 (p. 6), 3 (p. 8) and 5 (slides 6–10). Course slides, n.d.\nInitial concept: Social Interactive.pptx, slides 5, 7–8. User-supplied concept deck, n.d.',25)
  text(80,490,1380,'Project media + implementation',31,NAVY,True);text(80,545,1380,'Prototype screens: public-map-project repository, 2026; controlled browser captures.\nMap / avatars: supplied project assets; original attribution pending in source register.\nActivity icons / fish cut-outs: AI-assisted media; see the accompanying source register.\np5.js • Firebase Realtime Database • GitHub Pages • Web Audio API.',25)
  text(80,725,1380,'Supporting credits and AI-use details accompany this presentation.',22,MUTED)
 else:
  head('Appendix / not part of the timed talk','Before you submit');items=['Fill in your name; check the verified theory mapping and initial-concept comparison.','Record and insert a real 1–2 minute showcase film.','Rehearse a 5–7 minute presentation including video and live demo.','Prepare the separate development PDF: maximum 15 pages, excluding title / references.','Include coversheet, requested MP4 / SharePoint link and tutor access.','Confirm AI approval, declaration and media attribution; check milestone completion.']
  for j,t in enumerate(items):text(85,235+j*85,1410,f'{j+1:02d}  {t}',29,NAVY if j<3 else INK)
 notes=f"SLIDE {page_index+1} | {meta['time']}\n\nENGLISH SCRIPT\n{meta['en']}\n\n中文理解版\n{meta['zh']}\n\nPRESENTER CUE\n{meta['cue']}"
 sl.notes_slide.notes_text_frame.text=notes
 if i not in [0,10,11]:text(80,820,1400,'Images: prototype assets / controlled browser screenshots, project repository (2026).',14,MUTED)
 path=O/f'slide-{page_index+1:02d}.png';canvas.convert('RGB').save(path);canvases.append(canvas.convert('RGB'));manifest.append({'slide':page_index+1,**meta})
prs.save(O/'Shared-Campus-A2.pptx')
canvases[0].save(O/'Shared-Campus-A2.pdf',save_all=True,append_images=canvases[1:],resolution=120)
(O/'presentation-content.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
# Script, bilingual explanation and timings.
doc=Document();normal=doc.styles['Normal'];normal.font.name='Arial';normal.font.size=DPt(11)
doc.add_heading('Shared Campus — Presentation Script',0);doc.add_paragraph('GRDE2053 • A2 Exhibition Showcase • Draft for review • English delivery with Chinese explanation')
doc.add_paragraph('Target runtime: 6:45 including a 75-second showcase video and a 55-second live demonstration. Do not read the Chinese paragraphs aloud. Teacher documents require both video and live demonstration. The theory mapping is verified against the supplied course slides. Insert your real showcase video and name before presenting.')
for i,m in enumerate(slides[:-1]):
 doc.add_heading(f"{i+1:02d}  {m['title']}  |  {m['time']}",1)
 doc.add_paragraph(m['en']);doc.add_paragraph('中文理解：'+m['zh']);doc.add_paragraph('操作提示：'+m['cue'])
doc.add_heading('Timing and rehearsal',1);doc.add_paragraph('Use a timer. Aim for clear conversational delivery rather than reading every word. If the full rehearsal exceeds7 minutes, shorten technical detail and reflection; keep the required video, live demonstration and course-principle reflection. Rehearse using your actual phones and network.')
doc.add_heading('Important evidence boundaries',1);doc.add_paragraph('Development feedback and user confirmations in the project log are not a formal research study. Do not invent participant numbers, interview quotes, tutor feedback or A1 claims. The six principles are verified in Session 2 p.6 and Session 3 p.8. The initial concept comparison uses Social Interactive.pptx; confirm whether it is the actual A1 submission. Slide7 still needs real showcase footage. AI-assisted outputs must be disclosed according to tutor-approved unit rules.')
doc.save(O/'Speech-EN-with-Chinese-notes.docx')
md='# Shared Campus — 逐页演讲稿\n\n目标6:45，包含75秒视频和55秒现场演示。中文仅供理解，不要两种语言各读一遍。\n\n'
for i,m in enumerate(slides[:-1]):md+=f"## {i+1}. {m['title']} | {m['time']}\n\n{m['en']}\n\n**中文理解：** {m['zh']}\n\n**操作提示：** {m['cue']}\n\n"
(O/'Speech-EN-ZH.md').write_text(md)
print('Created editable PPTX, matching PDF, bilingual DOCX and15 slide previews.')
print('Spoken English words:',sum(len(re.findall(r"\b[\w’–-]+\b",m['en'])) for m in slides))
