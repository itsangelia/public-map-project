from pathlib import Path
from docx import Document
from docx.shared import Pt
import json,csv,zipfile,re,shutil
O=Path('/workspace/development-documentation-academic');D=json.loads((O/'page-content.json').read_text())
# Teacher requirement/evidence map is a preparation aid, not a substitute for visual evidence.
requirements=[('Brief p.3: concept development and deviations','1,12','Initial concept beside final map;identify retained and changed elements.'),('Brief p.3: technical experimentation and problem solving','2–4,8–9','System diagram,selected implementation details,and actual before/after records.'),('Brief p.3: feedback and responsive changes','3–10','Recorded feedback tied to a revision and re-check;do not relabel self-testing as audience research.'),('Brief p.3: visual process and participant planning','4–9,12','Phone layout,visual options,activity flow and exhibition arrangement.'),('Brief p.3: six Social UX principles','11','Six accurate names linked to project evidence and course sources;acknowledge unresolved barriers.'),('Brief p.3: AI use and adaptation','5,13','Specific assisted assets/code/writing uses and subsequent selection/adaptation.'),('Brief pp.3,5: image/research attribution','All figures + references','Figure numbers,captions and sources;complete original media credits and unit citation style.'),('Brief p.5: PDF maximum15 content pages','13 proposed content pages','Cover and reference list excluded;check actual page count after image layout.')]
notes=Document();notes.add_heading('A2正文修订与排版说明',0);notes.add_paragraph('依据本轮重新上传的A2_Brief.pdf第2–5页及A2评分表重新核对。本文件供准备使用，不是正文的一部分。')
notes.add_heading('本轮的写作调整',1)
for s in ['减少“I wanted / I asked / I liked”式口头叙述，保留有作者贡献意义的第一人称。','每页围绕设计条件、判断依据、实施变化和结果展开，避免写成操作流水账。','保留35%手机视角、20%减速、钓鱼容错、8秒让位等具体证据，避免抽象赞美。','区分可观察的功能变化、开发者主观反馈和仍待验证的社会效果。','视觉页补充氛围/体验设计；六原则页保留课程理论来源；A1比较不把自由探索和合作写成A2新构想。']:notes.add_paragraph(s)
notes.add_heading('老师要求与证据位置',1);t=notes.add_table(rows=1,cols=3);t.style='Light Shading Accent 1'
for c,v in zip(t.rows[0].cells,['要求','正文页','应配证据']):c.text=v
for r in requirements:
 for c,v in zip(t.add_row().cells,r):c.text=v
notes.add_heading('配图与引用',1)
for p in D:
 notes.add_heading(f"第{p['page']}页：{p['title']}",2);notes.add_paragraph(p['zh']);notes.add_paragraph(p['evidence']);notes.add_paragraph(p['caption']);notes.add_paragraph('来源依据：'+p['source'])
notes.add_heading('提交前必须核实的内容',1)
for s in ['Social Interactive.pptx与实际A1提交版是否一致。','真实前后截图、复测证据和测试时间；没有记录的结果不补写。','正式观众反馈尚未提供，正文仅写开发者设备检查；如已有实际观众反馈，可用真实记录补充第10页。','原地图、角色与游戏参考图的原始来源/许可，具体AI素材及工具记录。','unit outline指定的引用格式、导师对AI使用的批准及正式声明。第13页待核实状态在提交前须按事实更新。','最终PDF必须是视觉过程文档。不要把13页文字直接导出并当作图像证据已齐全的成品。']:notes.add_paragraph('□ '+s)
notes.save(O/'Writing-and-evidence-guide-ZH.docx')
checks=[['Observed issue','Documented response','Evidence boundary'],['Other players absent on a second phone/display','Review of joining/shared presence;later device confirmations','Recorded development checks;not proof of universal reliability'],['Small character/excess mobile whitespace','35% closer view,20% lower movement speed,overlaid controls','Selected through device feedback,not a measured optimal setting'],['Ingredient changes visually weak','More distinct ingredient rendering within the selected style','Use actual comparison states to support the claim'],['Jerky fishing movement and excessive difficulty','Movement refinement plus separate tolerance/timing/progress adjustments','Do not treat easier settings as a network fix'],['iPhone keys changed colour without sound','Permissions and browser audio activation addressed','Sound confirmed in device feedback;not all-device certification']]
with (O/'Testing-evidence-table.csv').open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(checks)
shutil.copy('/workspace/development-documentation/References-and-media-working-list.md',O/'References-and-media-working-list.md')
(O/'READ-ME-ZH.md').write_text('''# Development Documentation — 正式文体修订稿\n\n本版按2026年10月26日重传的brief和rubric复核，正文13页主题、约2100英文词。封面与参考文献另排；教师上限为正文15页。\n\n- Development-content-EN.docx：可排版的英文正文。\n- Development-content-EN-ZH.docx：同版英文＋中文说明、配图、图注和来源依据。\n- Writing-and-evidence-guide-ZH.docx：写作调整、教师要求映射与证据安排。\n- Testing-evidence-table.csv：基于记录的检查表，可编辑排版。\n- References-and-media-working-list.md：参考来源工作表，格式/授权仍需核实。\n\n本轮只调整文档，不修改正式地图。\n\n这份材料是文字内容与排版依据。完成真实截图、图片归属、引用、AI批准/声明和相关状态核实后，再导出提交PDF。新上传的教师文件未转载进材料包。\n''')
# Verify delivered language, traceability and integrity.
assert len(D)==13
for p in D:
 assert all(p[k] for k in ['title','en','zh','evidence','caption','source'])
 assert not any(s in p['en'] for s in ['I asked for','I liked','I wondered','AI-generated appearance','The evidence should show'])
for name in ['Development-content-EN.docx','Development-content-EN-ZH.docx']:
 d=Document(O/name);paras=[p.text for p in d.paragraphs]
 for p in D:
  for paragraph in p['en'].split('\n\n'):assert paragraph in paras
 assert zipfile.ZipFile(O/name).testzip() is None
for s in ['Sensory Feedback','Embodied Interaction','Social Dynamics','Resilience & Motivation','Proactivity & Clarity','Ethics & Inclusion']:assert s in D[10]['en']
with zipfile.ZipFile(O/'Development-content-package.zip','w',zipfile.ZIP_DEFLATED) as z:
 for f in O.iterdir():
  if f.suffix in ['.docx','.md','.csv','.json']:z.write(f,f.name)
print('Verified13 topics,English/bilingual alignment,and DOCX integrity.')
