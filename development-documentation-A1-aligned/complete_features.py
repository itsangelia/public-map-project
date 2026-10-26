from pathlib import Path
import json,csv,zipfile,re
from docx import Document
from docx.shared import Inches,Pt
O=Path('/workspace/development-documentation-A1-aligned');D=json.loads((O/'bilingual-content.json').read_text())
def replace(n,pairs):
 p=D[n-1];p['en']='\n\n'.join(x[0] for x in pairs);p['zh_full']='\n\n'.join(x[1] for x in pairs)
def append(n,idx,en,zh):
 p=D[n-1];a=p['en'].split('\n\n');b=p['zh_full'].split('\n\n');a[idx]+=' '+en;b[idx]+=zh;p['en']='\n\n'.join(a);p['zh_full']='\n\n'.join(b)
append(2,1,'The join page accepts a nickname and one of three avatars: male, female or shopping bag. Players enter the same room without creating a permanent account.','加入页面支持填写昵称，并从男性角色、女性角色或购物袋三种形象中选择。参与者无需创建永久账号，即可进入同一房间。')
append(3,2,'Connection and player-count information assist these checks. Reconnection logic republishes the player’s latest known position rather than treating a network return as a new shared character.','连接状态与玩家数量信息用于辅助检查。重新连接时，系统重新发布玩家最近的位置，而不是将网络恢复当作另一个共享角色。')
append(4,1,'A virtual joystick supports touch movement, the camera follows the local character and overhead nicknames identify players. Chat remains accessible from a separate phone control.','虚拟摇杆支持触控移动，镜头跟随本地角色，角色头顶的昵称用于识别玩家。手机上还保留独立的聊天入口。')
append(5,2,'The wishing tree, drink shop, fishing pond and piano each use approach cues, clickable entrances and a route back to the map.','许愿树、饮品店、钓鱼池和钢琴均配有靠近提示、可点击入口以及返回地图的方式。')
replace(6,[
('Chat supports immediate communication in the shared map. Players open the phone chat control and type a message, which is received by other players and appears above the sender’s character. The bubble disappears after five seconds, preventing old messages from remaining over the scene.','聊天支持共享地图中的即时交流。玩家通过手机聊天入口输入文字，其他玩家能够收到消息，文字显示在发送者角色头顶。气泡在五秒后消失，避免旧消息持续覆盖场景。'),
('The central wishing tree offers a persistent contribution. Its proximity cue invites players to open the tree interface, write a wish and select “Hang wish”. The composer retains the player’s nickname. The shared tree displays hanging notes, the total number of wishes and page navigation for browsing additional notes.','中央许愿树支持保留贡献。靠近提示引导玩家打开树的界面，写下愿望并选择“Hang wish”。编辑界面保留玩家昵称。共享树显示悬挂的纸条、愿望总数，并提供翻页功能以浏览更多纸条。'),
('Opening a note reveals the full wish, nickname and date within the framed-paper layout. Short previews and gentle movement encourage discovery. Earlier previews exposed too much text, so they were shortened to a brief opening of up to approximately five English words, leaving the complete message inside the note.','打开纸条后，可在带边框的纸张布局中查看完整愿望、昵称和日期。简短预览与轻微摆动鼓励发现。早期预览显示了过多文字，因此缩短为最多约五个英文词的开头，将完整内容保留在纸条内。'),
('The “May it come true” button adds a blessing and updates the shared count. Each player can bless a particular wish once. This provides acknowledgement without requiring another written message. Closing the note or leaving the tree returns participants to browsing or the map.','“May it come true”按钮用于送上祝福，并更新共享数量。每位玩家对同一愿望只能祝福一次。这让参与者无需另写留言，也能作出回应。关闭纸条或离开许愿树后，可以继续浏览或返回地图。')])
replace(7,[
('The drink activity develops A1’s proposal for short shared play through a two-player matching interface with nicknames and avatar portraits. The pair alternate across four stages: base, flavour, topping and decoration. Each stage has three options, producing 81 possible recipe combinations.','饮品活动通过显示昵称和角色头像的双人匹配界面，发展了A1关于简短共享玩乐的构想。两人依次在基底、风味、配料和装饰四个阶段轮流选择，每个阶段有三个选项，共可形成81种配方组合。'),
('The bases are sparkling water, milk tea and slush; flavours are strawberry, matcha and mango; toppings are pearls, coconut jelly and red bean; decorations are milk foam, ice cream and a lemon slice. Each selection changes the drink, and a “Your turn” notification identifies who can act next.','基底包括气泡水、奶茶和冰沙；风味包括草莓、抹茶和芒果；配料包括珍珠、椰果和红豆；装饰包括奶盖、冰激凌和柠檬片。每次选择都会改变饮品，“Your turn”提示标明接下来可以操作的玩家。'),
('Early feedback indicated that ingredient changes were not visually distinct enough. Revisions strengthened these differences while retaining the selected colour palette and upright cards. Cup logos, straws and the bottom shadow were removed from the visual treatment.','早期反馈指出，原料变化的视觉区别不够明显。修改增强了这些差异，同时保留已选配色和端正的选择卡片。视觉呈现中移除了杯身标志、吸管及底部阴影。'),
('Both players receive the same completed drink. The result screen enlarges it, displays its name and provides “Back to map”, without repeating the full selection list. The shared outcome records contributions from both participants and creates a reason to cooperate without competition.','两位玩家收到同一杯完成的饮品。结果界面将其放大、显示名称并提供“Back to map”，不再重复列出全部选择。共同结果体现双方贡献，为合作提供理由，而不需要竞争。')])
replace(8,[
('Fishing begins with two-player matching, nickname and avatar identification, followed by a countdown. One participant controls vertical movement and the other horizontal movement. Both influence one float, which must align with moving ripples to increase shared progress. Neither controls both axes independently.','钓鱼从双人匹配开始，通过昵称与头像识别双方，之后进入倒计时。一人控制上下，另一人控制左右。双方共同影响一个浮标，使其对准会换位的涟漪来增加共享进度，任何一人都不能独立控制两个轴向。'),
('Development feedback identified jerky movement and excessive difficulty as separate issues. Movement was refined for visual continuity. The target range was increased by 10%, target changes extended to 3.5-second intervals and progress loss reduced to 4% per second. Progress starts at 20%, increases by 10% per second when aligned and ends in failure at zero; reaching 100% completes the catch.','开发反馈将移动卡顿和难度过高识别为不同问题。移动表现经过调整，以提高视觉连续性。判定范围扩大10%，目标换位间隔延长至3.5秒，进度下降降低至每秒4%。进度从20%开始，对准时每秒增加10%，降至零则失败；达到100%完成捕获。'),
('Ripples replaced the visible fish, and the float appears partly submerged. Successful completion gives both players the same randomly revealed catch, with its image and name: Lotus Fish, Goofy Toothy, Bigmouth or Purrmaid. Failure shows Smug Finn with “The fish got away!”, without an additional escape animation. Retry and return controls support another attempt or a return to the map.','涟漪取代可见的鱼，浮标呈现为部分浸入水中的状态。成功后，双方获得相同的随机捕获结果，显示鱼的图像与名称：Lotus Fish、Goofy Toothy、Bigmouth或Purrmaid。失败时显示Smug Finn和“The fish got away!”，不增加额外逃走动画。重试与返回控件支持再次尝试或回到地图。'),
('These revisions retain continuous cooperation while making the task more forgiving. Difficulty adjustment remains distinct from synchronisation work: easier settings cannot resolve network or rendering problems by themselves.','这些修改在保留持续合作的同时提高容错。难度调整仍与同步处理分开：更容易的参数无法单独解决网络或渲染问题。')])
replace(9,[
('The piano is opened as an on-map panel with 21 numbered keys arranged in three rows. “Start playing” claims the instrument for one performer. Other players can listen across the map after enabling sound, with volume responding to distance from the piano. A sound control enables playback or mutes it.','钢琴以地图内面板打开，包含三排共21个数字音键。“Start playing”让一名演奏者占用乐器。其他玩家开启声音后可以在地图中聆听，音量随与钢琴的距离变化。声音控件用于开启播放或静音。'),
('Selecting an occupied piano shows the performer’s nickname rather than a duplicate keyboard. Listeners do not receive synchronised key highlights. The map entrance responds during performance with movement and musical-note effects, making the activity visible to others.','点击被占用的钢琴时，显示演奏者昵称，而不是另一套键盘。听众不会接收同步的按键高亮。地图入口在演奏时产生动作和音符效果，让其他人注意到正在进行的活动。'),
('An eight-second period without a note releases the seat; further notes refresh that period. Players can also finish playing or close the panel. The controls remain separated from the exit to reduce accidental selection. These mechanisms clarify occupancy and prevent an idle session from retaining the shared instrument.','八秒内没有按下音符会释放座位，继续弹奏则重新计算这一时段。玩家也可以主动结束弹奏或关闭面板。键盘与退出控件保持间距，以减少误触。这些机制明确占用状态，并避免闲置会话持续占据共享乐器。'),
('Technical checks exposed a separate audio problem: on an iPhone, keys changed colour without producing sound. Database permissions and browser audio activation required separate attention. Following these changes, device feedback confirmed audible playback. The issue demonstrated why a visible response could not establish that the full interaction was functioning.','技术检查还发现了独立的音频问题：iPhone上的按键会变色，却没有声音。数据库权限与浏览器音频激活需要分别处理。调整后，设备反馈确认能够听到声音。这说明，可见回应不能证明整个交互已正常运行。')])
# Keep original author guidance; the bilingual body now contains the complete feature description.
for p in D:assert len(p['en'].split('\n\n'))==len(p['zh_full'].split('\n\n'))
(O/'bilingual-content.json').write_text(json.dumps(D,ensure_ascii=False,indent=2))
(O/'page-content.json').write_text(json.dumps(D,ensure_ascii=False,indent=2))
for kind,name in [('en','Development-content-EN.docx'),('zh','Development-content-ZH.docx'),('parallel','Development-content-BILINGUAL.docx')]:
 d=Document();d.sections[0].top_margin=d.sections[0].bottom_margin=Inches(.7);d.styles['Normal'].font.name='Arial';d.styles['Normal'].font.size=Pt(10.5)
 for i,p in enumerate(D):
  if i:d.add_page_break()
  d.add_heading(f"{p['page']:02d}  "+(p['title'] if kind=='en' else p['title_zh']),0)
  if kind=='parallel':d.add_paragraph(p['title'],style='Subtitle')
  for en,zh in zip(p['en'].split('\n\n'),p['zh_full'].split('\n\n')):
   if kind in ['en','parallel']:d.add_paragraph(en)
   if kind in ['zh','parallel']:d.add_paragraph(zh)
  d.add_paragraph('Source basis / 来源依据: '+p['source'],style='Caption')
 d.save(O/name)
# Replace the older summary-only bilingual filename too, to avoid giving two different interpretations.
import shutil
shutil.copyfile(O/'Development-content-BILINGUAL.docx',O/'Development-content-EN-ZH.docx')
for kind,name in [('en','Development-content-EN.md'),('parallel','Development-content-BILINGUAL.md')]:
 md='# Shared Campus Map — Development Documentation\n\n'
 for p in D:
  md+=f"## {p['page']:02d}. {p['title']}"+('｜'+p['title_zh'] if kind=='parallel' else '')+'\n\n'
  for en,zh in zip(p['en'].split('\n\n'),p['zh_full'].split('\n\n')):md+=en+'\n\n'+(zh+'\n\n' if kind=='parallel' else '')
 (O/name).write_text(md)
F=[
('加入与共享地图','公共大屏与扫码入口；浏览器加入；昵称；男性/女性/购物袋三角色；同房间独立玩家；多人可见与位置同步','QR entry;browser joining;nickname;three avatars;independent players;shared-room presence and positions','2,3'),
('状态与恢复','连接状态和人数；独立连接；重新连接后发布最近位置','Connection/player count;independent connections;latest-position republishing after reconnect','3'),
('手机操作','竖屏跟随；地图铺满；35%视角放大；20%减速；摇杆；头顶昵称；聊天入口','Portrait follow camera;full map viewport;35% zoom;20% slower movement;joystick;overhead nicknames;chat control','4'),
('活动入口','许愿树/饮品/钓鱼/钢琴图标；靠近效果；点击进入；退出/返回地图','Four activity icons;proximity cues;clickable entries;exit/return routes','5'),
('聊天','输入文字；其他玩家接收；发送者头顶气泡；5秒消失','Typed messages;other-player reception;sender head bubbles;five-second expiry','6'),
('许愿树','中央树入口；昵称保留；写愿望；Hang wish；共享纸条；愿望总数与翻页','Central tree;retained nickname;write/hang;shared notes;wish count and pagination','6'),
('愿望阅读与祝福','短预览；轻微摆动；完整内容/昵称/日期；May it come true；共享祝福数；同一玩家一次祝福','Short preview;sway;full text/nickname/date;blessing button;shared count;one blessing per player per wish','6'),
('饮品匹配与轮次','双人匹配；昵称头像；四阶段轮流；每阶段三选一；Your turn','Two-player matching;names/portraits;four alternating stages;three options each;turn notification','7'),
('饮品选择','气泡水/奶茶/冰沙；草莓/抹茶/芒果；珍珠/椰果/红豆；奶盖/冰激凌/柠檬片；81组合','Three bases,flavours,toppings and finishes;81 combinations','7'),
('饮品反馈与结算','选择改变外观；同一结果；命名；成品放大；Back to map；去掉重复选择清单','Visible ingredient changes;shared named result;enlarged drink;return;no repeated choice list','7'),
('钓鱼匹配与操作','双人匹配昵称头像；倒计时；上下/左右分工；共享浮标；半浸浮标和换位涟漪','Matching names/portraits;countdown;split axes;shared submerged float;moving ripple target','8'),
('钓鱼进度','初始20%；对准+10%/秒；偏离−4%/秒；0失败/100%成功；范围扩大10%；3.5秒换位','Start20%;gain10%/s;loss4%/s;zero failure/full success;10% larger tolerance;3.5s target interval','8'),
('钓鱼结果','双方相同随机捕获；四鱼图像与名称；Smug Finn失败图；鱼逃走了；重试与返回','Shared random catch;four named fish;static escape fish;failure text;retry/return','8'),
('钢琴操作与声音','地图内21键三排；Start playing；单演奏者；全地图听众开启声音后聆听；距离音量；静音开关','On-map21-key panel;claim;one performer;enabled map-wide listening;distance volume;mute control','9'),
('钢琴占用与让位','演奏者昵称提示；听众不显示键盘与按键同步；演奏中图标/音符；8秒无音符释放；弹奏刷新计时；主动结束/关闭','Busy nickname notice;no listener keyboard/highlights;playing icon/notes;eight-second idle release;renewal;manual finish/close','9'),
('维护与验证','多人/消息/共同结果/声音检查；手机视角反馈；活动返回；稳定性与展览检查边界','Presence/message/shared result/audio checks;layout feedback;return paths;limits of verification','10,12')]
with (O/'Complete-feature-coverage.csv').open('w',encoding='utf-8-sig',newline='') as f:
 w=csv.writer(f);w.writerow(['功能组','中文功能','English coverage','正文页']);w.writerows(F)
d=Document();d.add_heading('完整功能覆盖核对 / Feature coverage',0);d.add_paragraph('本清单用于核对13页正文，不额外占正文页。只记录已实现功能；已删除的路灯活动、未实现的图书馆/背景音乐/代币不列为成品功能。检查了源代码与开发记录，不代表本轮重新完成了跨设备现场测试。')
for title,zh,en,p in F:
 d.add_heading(title+' — 第'+p+'页',2);d.add_paragraph(zh);d.add_paragraph(en)
d.save(O/'Complete-feature-coverage-EN-ZH.docx')
# Coverage checks and full bilingual consistency.
body='\n'.join(p['en'] for p in D)
for s in ['five seconds','Hang wish','page navigation','shared count','once','81','milk foam','countdown','Lotus Fish','Purrmaid','Smug Finn','Retry','21','eight-second','distance','mutes','nickname','joystick']:assert s in body,s
for kind,name in [('en','Development-content-EN.docx'),('zh','Development-content-ZH.docx'),('parallel','Development-content-BILINGUAL.docx')]:
 paras=[p.text for p in Document(O/name).paragraphs]
 for p in D:
  if kind!='zh':
   for t in p['en'].split('\n\n'):assert t in paras
  if kind!='en':
   for t in p['zh_full'].split('\n\n'):assert t in paras
 assert zipfile.ZipFile(O/name).testzip() is None
with zipfile.ZipFile(O/'Bilingual-complete-features-package.zip','w',zipfile.ZIP_DEFLATED) as z:
 for name in ['Development-content-BILINGUAL.docx','Development-content-EN.docx','Development-content-ZH.docx','Development-content-BILINGUAL.md','bilingual-content.json','Complete-feature-coverage-EN-ZH.docx','Complete-feature-coverage.csv','A1-alignment-and-evidence-guide-ZH.docx','References-and-media-working-list.md']:z.write(O/name,name)
print('Verified all13 page topics,complete translations and16 feature groups. English words:',sum(len(re.findall(r"\b[\w’–-]+\b",p['en'])) for p in D))
